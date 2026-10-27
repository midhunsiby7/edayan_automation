from __future__ import annotations

import hashlib
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from src.export.manual_review_exporter import ManualReviewExporter
from src.models import ProcessedRow, RecordType, TransactionType


REVIEW_SCHEMA_VERSION = "1.0"
REVIEW_SHEET = "Physical_Column_Review"
METADATA_SHEET = "Review_Metadata"
HUMAN_DIRECTIONS = ("", "Receipt", "Payment", "Ambiguous")
REVIEW_STATUSES = ("UNREVIEWED", "HUMAN_VERIFIED", "AMBIGUOUS", "NOT_REQUIRED")

REVIEW_COLUMNS = [
    "Source Image",
    "Source Page",
    "Source Row",
    "Record Type",
    "Date",
    "Voucher/Receipt Number",
    "Original Raw Amount",
    "Raw Amount Column 1",
    "Raw Amount Column 2",
    "Original Detected Amount Column",
    "Original Model Classification",
    "Human-Verified Direction",
    "Review Status",
    "Correction Reason",
    "Head Name",
    "Subhead",
    "Amount",
    "Narration",
    "Priority Review",
]

def source_key(row: Any) -> tuple[str, str, int]:
    return (str(row.source_image), str(row.source_page), int(row.source_row))


def key_as_text(key: tuple[str, str, int]) -> str:
    return f"{key[0]}|{key[1]}|{key[2]}"


def extraction_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _validate_unique(rows: Sequence[Any], label: str) -> dict[tuple[str, str, int], Any]:
    keys = [source_key(row) for row in rows]
    duplicates = sorted(key for key, count in Counter(keys).items() if count > 1)
    if duplicates:
        raise ValueError(f"{label} contains duplicate source identifiers: {duplicates}")
    return {source_key(row): row for row in rows}


def _priority_keys() -> set[tuple[str, str, int]]:
    return {
        ("page_001.jpg", "1", row)
        for row in (5, 6, 7, 8, 10, 14, 15, 22, 24, 25)
    } | {
        ("page_002.jpg", "2", row)
        for row in (3, 4)
    }


def build_review_records(
    processed_rows: Sequence[ProcessedRow],
    extraction_path: Path,
) -> list[dict[str, Any]]:
    processed = _validate_unique(processed_rows, "processed rows")
    extraction_hash = extraction_sha256(extraction_path)
    priority = _priority_keys()
    records: list[dict[str, Any]] = []

    for key, row in sorted(processed.items(), key=lambda item: item[0]):
        informational = row.record_type != RecordType.TRANSACTION
        records.append(
            {
                "source_image": key[0],
                "source_page": key[1],
                "source_row": key[2],
                "record_type": row.record_type.value,
                "date": row.resolved_date,
                "voucher_number": row.raw_voucher_number,
                "raw_amount": row.raw_amount,
                "raw_amount_col1": row.raw_amount_col1,
                "raw_amount_col2": row.raw_amount_col2,
                "detected_amount_column": row.detected_amount_column,
                "model_classification": row.transaction_type.value,
                "human_verified_direction": None,
                "review_status": "NOT_REQUIRED" if informational else "UNREVIEWED",
                "correction_reason": "",
                "head_name": row.accounting_head_name,
                "subhead": row.subhead or "",
                "amount": row.receipt_amount if row.transaction_type == TransactionType.RECEIPT else row.payment_amount,
                "narration": row.narration_english or "",
                "priority_review": key in priority,
                "source_extraction_sha256": extraction_hash,
            }
        )
    return records


def write_review_workbook(
    records: Sequence[Mapping[str, Any]],
    output_path: Path,
    *,
    source_path: Path,
    source_hash: str,
) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = REVIEW_SHEET
    sheet.append(REVIEW_COLUMNS)

    header_fill = PatternFill("solid", fgColor="1F4E79")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for record in records:
        sheet.append(
            [
                record["source_image"],
                record["source_page"],
                record["source_row"],
                record["record_type"],
                record["date"],
                record["voucher_number"],
                record["raw_amount"],
                record["raw_amount_col1"],
                record["raw_amount_col2"],
                record["detected_amount_column"],
                record["model_classification"],
                record["human_verified_direction"] or "",
                record["review_status"],
                record["correction_reason"],
                record["head_name"],
                record["subhead"],
                record["amount"],
                record["narration"],
                "PRIORITY_REVIEW" if record["priority_review"] else "",
            ]
        )

    direction_validation = DataValidation(
        type="list",
        formula1='"Receipt,Payment,Ambiguous"',
        allow_blank=True,
    )
    status_validation = DataValidation(
        type="list",
        formula1='"UNREVIEWED,HUMAN_VERIFIED,AMBIGUOUS,NOT_REQUIRED"',
        allow_blank=False,
    )
    sheet.add_data_validation(direction_validation)
    sheet.add_data_validation(status_validation)
    direction_validation.add(f"L2:L{sheet.max_row}")
    status_validation.add(f"M2:M{sheet.max_row}")
    sheet.conditional_formatting.add(
        f"L2:M{sheet.max_row}",
        FormulaRule(formula=['OR($S2="PRIORITY_REVIEW",$M2="UNREVIEWED",$M2="AMBIGUOUS")'], fill=PatternFill("solid", fgColor="FFF3CD")),
    )

    widths = [18, 12, 10, 18, 14, 20, 18, 20, 20, 18, 24, 24, 18, 32, 42, 24, 14, 55, 18]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    metadata = workbook.create_sheet(METADATA_SHEET)
    metadata.sheet_state = "hidden"
    metadata_rows = [
        ("review_schema_version", REVIEW_SCHEMA_VERSION),
        ("source_path", str(source_path)),
        ("source_sha256", source_hash),
        ("generated_at_utc", datetime.now(timezone.utc).isoformat()),
        ("decision_policy", "Only HUMAN_VERIFIED Receipt/Payment decisions may be finalized."),
        ("row_bbox_policy", "Any source row box is locator context only, not amount-cell verification."),
    ]
    for row in metadata_rows:
        metadata.append(row)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def _metadata(workbook: openpyxl.Workbook) -> dict[str, str]:
    if METADATA_SHEET not in workbook.sheetnames:
        raise ValueError(f"Missing required hidden sheet: {METADATA_SHEET}")
    return {
        str(row[0].value): str(row[1].value)
        for row in workbook[METADATA_SHEET].iter_rows(min_row=1, max_col=2)
        if row[0].value is not None
    }


def _validate_decision(direction: str | None, status: str, informational: bool, key: tuple[str, str, int]) -> None:
    direction = direction or ""
    if status not in REVIEW_STATUSES:
        raise ValueError(f"{key_as_text(key)} has invalid review status {status!r}")
    if informational:
        if status != "NOT_REQUIRED" or direction:
            raise ValueError(f"{key_as_text(key)} informational record must be NOT_REQUIRED with no direction")
        return
    if status == "HUMAN_VERIFIED" and direction not in {"Receipt", "Payment"}:
        raise ValueError(f"{key_as_text(key)} HUMAN_VERIFIED requires Receipt or Payment")
    if status == "AMBIGUOUS" and direction != "Ambiguous":
        raise ValueError(f"{key_as_text(key)} AMBIGUOUS requires Ambiguous direction")
    if status == "UNREVIEWED" and direction:
        raise ValueError(f"{key_as_text(key)} UNREVIEWED cannot contain a direction")
    if status == "NOT_REQUIRED":
        raise ValueError(f"{key_as_text(key)} transaction cannot be NOT_REQUIRED")


def _validate_correction_reason(
    model_classification: str,
    direction: str | None,
    status: str,
    reason: str,
    key: tuple[str, str, int],
) -> None:
    if status != "HUMAN_VERIFIED" or direction not in {"Receipt", "Payment"}:
        return
    model_direction = {
        TransactionType.RECEIPT.value: "Receipt",
        TransactionType.PAYMENT.value: "Payment",
    }.get(model_classification)
    if model_direction != direction and not reason.strip():
        raise ValueError(
            f"{key_as_text(key)} requires a correction reason when human direction "
            "differs from the model proposal"
        )


def import_review_workbook(
    workbook_path: Path,
    extraction_path: Path,
    processed_rows: Sequence[ProcessedRow],
    sidecar_path: Path,
    *,
    reviewer: str | None = None,
) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(workbook_path, data_only=True)
    if REVIEW_SHEET not in workbook.sheetnames:
        raise ValueError(f"Missing review sheet {REVIEW_SHEET!r}")
    metadata = _metadata(workbook)
    current_hash = extraction_sha256(extraction_path)
    if metadata.get("source_sha256") != current_hash:
        raise ValueError("Review workbook extraction hash does not match current raw extraction")
    if metadata.get("review_schema_version") != REVIEW_SCHEMA_VERSION:
        raise ValueError("Unsupported physical-column review schema version")

    expected = _validate_unique(processed_rows, "processed rows")
    sheet = workbook[REVIEW_SHEET]
    headers = [cell.value for cell in sheet[1]]
    if headers != REVIEW_COLUMNS:
        raise ValueError("Review workbook columns do not match the current schema")
    seen: set[tuple[str, str, int]] = set()
    decisions: list[dict[str, Any]] = []
    positions = {name: index for index, name in enumerate(headers)}

    for values in sheet.iter_rows(min_row=2, values_only=True):
        if all(value is None for value in values):
            continue
        try:
            key = (
                str(values[positions["Source Image"]]),
                str(values[positions["Source Page"]]),
                int(values[positions["Source Row"]]),
            )
        except (TypeError, ValueError) as exc:
            raise ValueError(f"Invalid source identifier in review workbook: {values[:3]}") from exc
        if key in seen:
            raise ValueError(f"Duplicate source identifier in review workbook: {key_as_text(key)}")
        seen.add(key)
        if key not in expected:
            raise ValueError(f"Unexpected source identifier in review workbook: {key_as_text(key)}")

        row = expected[key]
        model_classification = str(values[positions["Original Model Classification"]])
        if model_classification != row.transaction_type.value:
            raise ValueError(f"Model classification changed for {key_as_text(key)}")
        direction = values[positions["Human-Verified Direction"]]
        status = str(values[positions["Review Status"]] or "")
        correction_reason = str(values[positions["Correction Reason"]] or "")
        _validate_decision(direction, status, row.record_type != RecordType.TRANSACTION, key)
        _validate_correction_reason(
            model_classification,
            direction,
            status,
            correction_reason,
            key,
        )
        decisions.append(
            {
                "source_image": key[0],
                "source_page": key[1],
                "source_row": key[2],
                "original_model_detected_amount_column": row.detected_amount_column,
                "original_model_classification": row.transaction_type.value,
                "human_verified_direction": direction or None,
                "review_status": status,
                "correction_reason": correction_reason,
                "reviewer": reviewer,
                "reviewed_at_utc": datetime.now(timezone.utc).isoformat()
                if status in {"HUMAN_VERIFIED", "AMBIGUOUS"}
                else None,
                "source_extraction_sha256": current_hash,
                "review_schema_version": REVIEW_SCHEMA_VERSION,
            }
        )

    missing = sorted(set(expected) - seen)
    if missing:
        raise ValueError(f"Review workbook is missing source identifiers: {[key_as_text(k) for k in missing]}")

    counts = Counter(decision["review_status"] for decision in decisions)
    payload = {
        "review_schema_version": REVIEW_SCHEMA_VERSION,
        "source_path": str(extraction_path),
        "source_sha256": current_hash,
        "imported_at_utc": datetime.now(timezone.utc).isoformat(),
        "reviewer": reviewer,
        "records": decisions,
        "counts": dict(counts),
    }
    sidecar_path.parent.mkdir(parents=True, exist_ok=True)
    sidecar_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload


def load_decisions(
    sidecar_path: Path,
    extraction_path: Path,
    processed_rows: Sequence[ProcessedRow],
) -> dict[tuple[str, str, int], dict[str, Any]]:
    payload = json.loads(sidecar_path.read_text(encoding="utf-8"))
    current_hash = extraction_sha256(extraction_path)
    if payload.get("review_schema_version") != REVIEW_SCHEMA_VERSION:
        raise ValueError("Unsupported physical-column review schema version")
    if payload.get("source_sha256") != current_hash:
        raise ValueError("Review sidecar extraction hash does not match current raw extraction")
    expected = _validate_unique(processed_rows, "processed rows")
    records = payload.get("records")
    if not isinstance(records, list):
        raise ValueError("Review sidecar records must be a list")
    result: dict[tuple[str, str, int], dict[str, Any]] = {}
    for record in records:
        key = (str(record["source_image"]), str(record["source_page"]), int(record["source_row"]))
        if key in result:
            raise ValueError(f"Duplicate source identifier in review sidecar: {key_as_text(key)}")
        if key not in expected:
            raise ValueError(f"Unexpected source identifier in review sidecar: {key_as_text(key)}")
        row = expected[key]
        if record.get("original_model_classification") != row.transaction_type.value:
            raise ValueError(f"Sidecar model classification mismatch for {key_as_text(key)}")
        if record.get("source_extraction_sha256") != current_hash:
            raise ValueError(f"Sidecar record hash mismatch for {key_as_text(key)}")
        _validate_decision(
            record.get("human_verified_direction"),
            str(record.get("review_status", "")),
            row.record_type != RecordType.TRANSACTION,
            key,
        )
        _validate_correction_reason(
            str(record.get("original_model_classification")),
            record.get("human_verified_direction"),
            str(record.get("review_status", "")),
            str(record.get("correction_reason", "")),
            key,
        )
        result[key] = record
    missing = sorted(set(expected) - set(result))
    if missing:
        raise ValueError(f"Review sidecar is missing source identifiers: {[key_as_text(k) for k in missing]}")
    return result


def build_reconciliation_report(
    processed_rows: Sequence[ProcessedRow],
    decisions: Mapping[tuple[str, str, int], Mapping[str, Any]],
) -> dict[str, Any]:
    _validate_unique(processed_rows, "processed rows")
    eligible = [row for row in processed_rows if row.record_type == RecordType.TRANSACTION]
    unresolved = [
        key_as_text(source_key(row))
        for row in eligible
        if decisions[source_key(row)].get("review_status") != "HUMAN_VERIFIED"
    ]
    verified = [
        row for row in eligible if decisions[source_key(row)].get("review_status") == "HUMAN_VERIFIED"
    ]
    receipt_keys = [
        key_as_text(source_key(row))
        for row in verified
        if decisions[source_key(row)].get("human_verified_direction") == "Receipt"
    ]
    payment_keys = [
        key_as_text(source_key(row))
        for row in verified
        if decisions[source_key(row)].get("human_verified_direction") == "Payment"
    ]
    ambiguous = [
        key_as_text(source_key(row))
        for row in eligible
        if decisions[source_key(row)].get("review_status") == "AMBIGUOUS"
    ]
    unreviewed = [
        key_as_text(source_key(row))
        for row in eligible
        if decisions[source_key(row)].get("review_status") == "UNREVIEWED"
    ]
    return {
        "source_record_count": len(processed_rows),
        "eligible_transaction_count": len(eligible),
        "informational_excluded_count": len(processed_rows) - len(eligible),
        "human_verified_count": len(verified),
        "receipt_count": len(receipt_keys),
        "payment_count": len(payment_keys),
        "ambiguous_count": len(ambiguous),
        "unreviewed_count": len(unreviewed),
        "receipt_source_keys": receipt_keys,
        "payment_source_keys": payment_keys,
        "ambiguous_source_keys": ambiguous,
        "unreviewed_source_keys": unreviewed,
        "unresolved_source_keys": unresolved,
        "finalization_allowed": not unresolved,
    }


def export_verified_workbook(
    processed_rows: Sequence[ProcessedRow],
    decisions: Mapping[tuple[str, str, int], Mapping[str, Any]],
    output_path: Path,
) -> Path:
    report = build_reconciliation_report(processed_rows, decisions)
    if not report["finalization_allowed"]:
        raise ValueError(
            "Finalization blocked: unresolved physical-column decisions remain "
            f"({len(report['unresolved_source_keys'])})"
        )
    directions: dict[tuple[str, str, int], TransactionType] = {}
    for row in processed_rows:
        if row.record_type != RecordType.TRANSACTION:
            continue
        direction = decisions[source_key(row)]["human_verified_direction"]
        directions[source_key(row)] = (
            TransactionType.RECEIPT if direction == "Receipt" else TransactionType.PAYMENT
        )
    return ManualReviewExporter().export_with_directions(processed_rows, directions, output_path)
