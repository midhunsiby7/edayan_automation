import logging
from pathlib import Path
from typing import List
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.models import ProcessedRow, RecordType, TransactionType

logger = logging.getLogger(__name__)

# Colors
HEADER_FILL_COLOR = "1F4E79"       # Deep Navy
HEADER_FONT_COLOR = "FFFFFF"       # White
REVIEW_WARN_FILL = "FFF3CD"        # Soft Amber / Yellow warning
ZEBRA_FILL = "F9FAFB"              # Light neutral alternate fill
BORDER_COLOR = "D9D9D9"            # Subtle light gray


class ExcelExporter:
    """Exports processed cashbook records to a multi-sheet, audit-ready Excel workbook."""

    STANDARD_COLUMNS = [
        ("Source Image", "source_image", 20),
        ("Page", "source_page", 8),
        ("Row", "source_row", 8),
        ("Raw Date", "raw_date", 12),
        ("Resolved Date", "resolved_date", 14),
        ("Acct Head No", "accounting_head_number", 14),
        ("Acct Head Name", "accounting_head_name", 25),
        ("Web Head Value", "website_head_value", 16),
        ("Web Head Label", "website_head_label", 30),
        ("Voucher No", "voucher_number", 12),
        ("Raw Narration (Malayalam)", "raw_narration", 35),
        ("English Narration", "english_narration", 30),
        ("Receipt Amount", "receipt_amount", 16),
        ("Payment Amount", "payment_amount", 16),
        ("Transaction Type", "transaction_type", 16),
        ("Record Type", "record_type", 18),
        ("Confidence", "extraction_confidence", 12),
        ("Review Required", "review_required", 15),
    ]

    AUDIT_COLUMNS = [
        ("Source Image", "source_image", 20),
        ("Page", "source_page", 8),
        ("Row", "source_row", 8),
        ("BBox [y1,x1,y2,x2]", "row_bbox", 22),
        ("Raw Date", "raw_date", 12),
        ("Resolved Date", "resolved_date", 14),
        ("Date Conf", "date_confidence", 10),
        ("Acct Head No", "accounting_head_number", 14),
        ("Acct Head Name", "accounting_head_name", 25),
        ("Web Head Value", "website_head_value", 16),
        ("Web Head Label", "website_head_label", 30),
        ("Head Conf", "head_confidence", 10),
        ("Voucher No", "voucher_number", 12),
        ("Voucher Conf", "voucher_confidence", 12),
        ("Raw Narration (Malayalam)", "raw_narration", 32),
        ("English Narration", "english_narration", 28),
        ("Narration Conf", "narration_confidence", 12),
        ("Raw Col 1 (Receipt Col)", "raw_amount_col1", 18),
        ("Raw Col 2 (Payment Col)", "raw_amount_col2", 18),
        ("Detected Col Index", "detected_amount_column", 16),
        ("Resolved Receipt", "receipt_amount", 16),
        ("Resolved Payment", "payment_amount", 16),
        ("Amount Conf", "amount_confidence", 12),
        ("Transaction Type", "transaction_type", 16),
        ("Record Type", "record_type", 18),
        ("Overall Conf", "extraction_confidence", 12),
        ("Review Required", "review_required", 15),
        ("Review Reasons / Warnings", "review_reasons", 45),
    ]

    def export(self, rows: List[ProcessedRow], output_path: Path) -> Path:
        """Export processed rows to an Excel workbook with Receipts, Payments, All_Transactions, and Extraction_Audit."""
        logger.info(f"Generating Excel workbook with {len(rows)} rows at {output_path}")
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remove default empty sheet

        # Filter subsets
        receipt_rows = [
            r for r in rows
            if r.transaction_type == TransactionType.RECEIPT
        ]
        payment_rows = [
            r for r in rows
            if r.transaction_type == TransactionType.PAYMENT
        ]

        # 1. Receipts Sheet
        self._write_sheet(wb, "Receipts", self.STANDARD_COLUMNS, receipt_rows)

        # 2. Payments Sheet
        self._write_sheet(wb, "Payments", self.STANDARD_COLUMNS, payment_rows)

        # 3. All_Transactions Sheet
        self._write_sheet(wb, "All_Transactions", self.STANDARD_COLUMNS, rows)

        # 4. Extraction_Audit Sheet
        self._write_audit_sheet(wb, "Extraction_Audit", self.AUDIT_COLUMNS, rows)

        # Save workbook
        output_path.parent.mkdir(parents=True, exist_ok=True)
        try:
            wb.save(output_path)
            logger.info(f"Successfully saved Excel file: {output_path}")
        except PermissionError:
            import time
            fallback_path = output_path.with_name(f"{output_path.stem}_{int(time.time())}.xlsx")
            logger.warning(
                f"Target file {output_path.name} is open in Excel or locked. "
                f"Saving to fallback path: {fallback_path.name}"
            )
            wb.save(fallback_path)
            output_path = fallback_path

        return output_path

    def _write_sheet(
        self,
        wb: openpyxl.Workbook,
        sheet_title: str,
        columns: list,
        data_rows: List[ProcessedRow],
    ):
        ws = wb.create_sheet(title=sheet_title)
        ws.views.sheetView[0].showGridLines = True

        header_font = Font(name="Calibri", size=11, bold=True, color=HEADER_FONT_COLOR)
        header_fill = PatternFill(start_color=HEADER_FILL_COLOR, end_color=HEADER_FILL_COLOR, fill_type="solid")
        thin_side = Side(border_style="thin", color=BORDER_COLOR)
        grid_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")

        # Write header
        for col_idx, (col_name, _, width) in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center
            cell.border = grid_border
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = max(width, len(col_name) + 3)

        ws.row_dimensions[1].height = 26

        # Write data rows
        warn_fill = PatternFill(start_color=REVIEW_WARN_FILL, end_color=REVIEW_WARN_FILL, fill_type="solid")
        zebra_fill = PatternFill(start_color=ZEBRA_FILL, end_color=ZEBRA_FILL, fill_type="solid")

        for r_idx, row_data in enumerate(data_rows, start=2):
            is_warn = row_data.review_required
            is_zebra = (r_idx % 2 == 0)

            for col_idx, (_, field_key, _) in enumerate(columns, start=1):
                val = getattr(row_data, field_key, None)

                # Format enums or lists
                if isinstance(val, (TransactionType, RecordType)):
                    val = val.value
                elif isinstance(val, list):
                    val = ", ".join(str(x) for x in val)

                cell = ws.cell(row=r_idx, column=col_idx, value=val)
                cell.border = grid_border

                if is_warn:
                    cell.fill = warn_fill
                elif is_zebra:
                    cell.fill = zebra_fill

                # Alignment and number formatting
                if field_key in ("receipt_amount", "payment_amount"):
                    cell.alignment = align_right
                    cell.number_format = "#,##0.00"
                elif field_key in ("source_page", "source_row", "resolved_date", "raw_date", "accounting_head_number", "website_head_value", "voucher_number", "review_required"):
                    cell.alignment = align_center
                elif field_key in ("extraction_confidence",):
                    cell.alignment = align_center
                    cell.number_format = "0.00"
                else:
                    cell.alignment = align_left

            ws.row_dimensions[r_idx].height = 20

    def _write_audit_sheet(
        self,
        wb: openpyxl.Workbook,
        sheet_title: str,
        columns: list,
        data_rows: List[ProcessedRow],
    ):
        ws = wb.create_sheet(title=sheet_title)
        ws.views.sheetView[0].showGridLines = True

        header_font = Font(name="Calibri", size=10, bold=True, color=HEADER_FONT_COLOR)
        audit_header_fill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
        thin_side = Side(border_style="thin", color=BORDER_COLOR)
        grid_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")

        for col_idx, (col_name, _, width) in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = audit_header_fill
            cell.alignment = align_center
            cell.border = grid_border
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = max(width, len(col_name) + 3)

        ws.row_dimensions[1].height = 26

        warn_fill = PatternFill(start_color=REVIEW_WARN_FILL, end_color=REVIEW_WARN_FILL, fill_type="solid")
        zebra_fill = PatternFill(start_color=ZEBRA_FILL, end_color=ZEBRA_FILL, fill_type="solid")

        for r_idx, row_data in enumerate(data_rows, start=2):
            is_warn = row_data.review_required
            is_zebra = (r_idx % 2 == 0)

            for col_idx, (_, field_key, _) in enumerate(columns, start=1):
                val = getattr(row_data, field_key, None)

                if isinstance(val, (TransactionType, RecordType)):
                    val = val.value
                elif isinstance(val, list):
                    val = "; ".join(str(x) for x in val)

                cell = ws.cell(row=r_idx, column=col_idx, value=val)
                cell.border = grid_border

                if is_warn:
                    cell.fill = warn_fill
                elif is_zebra:
                    cell.fill = zebra_fill

                if "amount" in field_key and isinstance(val, (int, float)):
                    cell.alignment = align_right
                    cell.number_format = "#,##0.00"
                elif "confidence" in field_key or "conf" in field_key:
                    cell.alignment = align_center
                    cell.number_format = "0.00"
                elif field_key in ("source_page", "source_row", "resolved_date", "raw_date", "accounting_head_number", "website_head_value", "voucher_number", "review_required", "detected_amount_column"):
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left

            ws.row_dimensions[r_idx].height = 20
