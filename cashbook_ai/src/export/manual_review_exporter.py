from pathlib import Path
from typing import Iterable, List

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from src.models import ProcessedRow, RecordType, TransactionType


MANUAL_REVIEW_COLUMNS = [
    "Date",
    "Receipt Number",
    "Amount",
    "Head Name",
    "Subhead",
    "Narration",
]


class ManualReviewExporter:
    """Exports transaction rows to the intentionally minimal monthly review workbook."""

    def export(self, rows: Iterable[ProcessedRow], output_path: Path) -> Path:
        rows = list(rows)
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)

        receipts = [
            row for row in rows
            if row.transaction_type == TransactionType.RECEIPT
            and row.record_type == RecordType.TRANSACTION
        ]
        payments = [
            row for row in rows
            if row.transaction_type == TransactionType.PAYMENT
            and row.record_type == RecordType.TRANSACTION
        ]

        self._write_sheet(workbook, "Receipts", receipts, "Receipt")
        self._write_sheet(workbook, "Payments", payments, "Payment")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
        return output_path

    def export_with_directions(
        self,
        rows: Iterable[ProcessedRow],
        directions: dict[tuple[str, str, int], TransactionType],
        output_path: Path,
        *,
        canonical_head_names: Iterable[str] | None = None,
        head_name_overrides: dict[tuple[str, str, int], str | None] | None = None,
    ) -> Path:
        """Export using explicit human-approved directions without mutating processed rows."""
        rows = list(rows)
        head_name_overrides = head_name_overrides or {}
        receipts = [
            row for row in rows
            if row.record_type == RecordType.TRANSACTION
            and directions.get((row.source_image, row.source_page, row.source_row)) == TransactionType.RECEIPT
        ]
        payments = [
            row for row in rows
            if row.record_type == RecordType.TRANSACTION
            and directions.get((row.source_image, row.source_page, row.source_row)) == TransactionType.PAYMENT
        ]

        workbook = Workbook()
        workbook.remove(workbook.active)
        self._write_sheet(
            workbook,
            "Receipts",
            receipts,
            "Receipt",
            head_name_overrides=head_name_overrides,
        )
        self._write_sheet(
            workbook,
            "Payments",
            payments,
            "Payment",
            head_name_overrides=head_name_overrides,
        )
        if canonical_head_names is not None:
            self._add_head_name_validation(workbook, list(canonical_head_names))
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
        return output_path

    def _write_sheet(
        self,
        workbook: openpyxl.Workbook,
        title: str,
        rows: List[ProcessedRow],
        amount_label: str,
        *,
        head_name_overrides: dict[tuple[str, str, int], str | None] | None = None,
    ) -> None:
        worksheet = workbook.create_sheet(title)
        worksheet.append(MANUAL_REVIEW_COLUMNS)

        header_fill = PatternFill("solid", fgColor="1F4E79")
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        head_name_overrides = head_name_overrides or {}
        unresolved_rows = []
        for row in rows:
            # A verified direction may override the model direction. In that
            # case the parsed amount is still the same source transaction
            # amount, so fall back to the other resolved side.
            if amount_label == "Receipt":
                amount = row.receipt_amount if row.receipt_amount is not None else row.payment_amount
            else:
                amount = row.payment_amount if row.payment_amount is not None else row.receipt_amount
            # The raw voucher is retained so unusual or uncertain source identifiers
            # are not silently rewritten by numeric normalization.
            receipt_number = row.raw_voucher_number
            key = (row.source_image, row.source_page, row.source_row)
            head_name = head_name_overrides.get(key, row.accounting_head_name)
            if not head_name:
                unresolved_rows.append(worksheet.max_row + 1)
            worksheet.append([
                row.resolved_date,
                receipt_number,
                amount,
                head_name or "",
                row.subhead or "",
                row.narration_english or "",
            ])

        widths = [14, 18, 14, 42, 24, 55]
        for index, width in enumerate(widths, start=1):
            worksheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        if unresolved_rows:
            for row_number in unresolved_rows:
                for column in (4,):
                    worksheet.cell(row_number, column).fill = PatternFill(
                        "solid",
                        fgColor="FFF3CD",
                    )

        for row in worksheet.iter_rows(min_row=2, min_col=3, max_col=3):
            row[0].number_format = "#,##0.00"

    def _add_head_name_validation(self, workbook: Workbook, canonical_names: List[str]) -> None:
        unique_names = list(dict.fromkeys(name.strip() for name in canonical_names if name and name.strip()))
        helper = workbook.create_sheet("__HeadNames")
        helper.sheet_state = "hidden"
        helper["A1"] = "Canonical human_final_name"
        for index, name in enumerate(unique_names, start=2):
            helper.cell(index, 1).value = name
        end_row = max(2, len(unique_names) + 1)
        workbook.defined_names.add(
            DefinedName(
                "CanonicalHeadNames",
                attr_text=f"'__HeadNames'!$A$2:$A${end_row}",
            )
        )
        validation = DataValidation(
            type="list",
            formula1="=CanonicalHeadNames",
            allow_blank=True,
        )
        validation.error = "Select a canonical head name from the verified master."
        validation.errorTitle = "Invalid head name"
        validation.prompt = "Choose a canonical human-verified head name."
        validation.promptTitle = "Head Name"
        for sheet_name in ("Receipts", "Payments"):
            sheet = workbook[sheet_name]
            sheet.add_data_validation(validation)
            validation.add(f"D2:D{max(sheet.max_row, 200)}")
