import csv
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

def main():
    draft_csv = Path("data/reference/accounting_heads_general_draft.csv")
    out_xlsx = Path("data/reference/accounting_heads_general_human_review.xlsx")
    
    # 1. Read the draft
    rows = []
    with open(draft_csv, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            rows.append(row)
            
    # 2. Create Workbook
    wb = openpyxl.Workbook()
    
    # --- Instructions Sheet ---
    ws_instr = wb.active
    ws_instr.title = "Review Instructions"
    
    instructions = [
        "Step 1: Open the original Head Master photograph.",
        "Step 2: Find the corresponding head number.",
        "Step 3: Compare the handwritten head name with extracted_head_name.",
        "Step 4: If the extracted value is correct: review_status = CORRECT | human_final_name = same name",
        "Step 5: If the extracted value is wrong: review_status = CORRECTED | human_final_name = the exact name the human confirms from the book",
        "Step 6: If the handwriting cannot be confidently determined: review_status = UNCERTAIN | human_final_name = blank | human_comment = explain the issue",
        "",
        "DO NOT automatically copy extracted_head_name into human_final_name as a verified value.",
        "After finishing, reply to the AI with: VERIFY HEAD MASTER"
    ]
    
    for instr in instructions:
        ws_instr.append([instr])
        
    ws_instr.column_dimensions['A'].width = 120
    
    # --- Human Review Sheet ---
    ws_review = wb.create_sheet("Human Review")
    
    headers = [
        "accounting_head_number",
        "extracted_head_name",
        "human_final_name",
        "review_status",
        "human_comment",
        "source_image",
        "source_page",
        "source_row"
    ]
    
    ws_review.append(headers)
    
    # Style header
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
    for col_idx in range(1, len(headers) + 1):
        cell = ws_review.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        
    # Add Data Validation for review_status
    dv = DataValidation(type="list", formula1='"CORRECT,CORRECTED,UNCERTAIN"', allow_blank=True)
    ws_review.add_data_validation(dv)
    
    for row_idx, row in enumerate(rows, start=2):
        ws_review.append([
            row["accounting_head_number"],
            row["accounting_head_name"],
            "", # human_final_name (blank initially)
            "", # review_status
            "", # human_comment
            row["source_image"],
            row["source_page"],
            row["source_row"]
        ])
        # add validation to the status cell
        dv.add(ws_review.cell(row=row_idx, column=4))
        
    # Freeze top row
    ws_review.freeze_panes = "A2"
    
    # Enable filters
    ws_review.auto_filter.ref = ws_review.dimensions
    
    # Widen columns
    ws_review.column_dimensions['A'].width = 25
    ws_review.column_dimensions['B'].width = 40
    ws_review.column_dimensions['C'].width = 40
    ws_review.column_dimensions['D'].width = 20
    ws_review.column_dimensions['E'].width = 40
    ws_review.column_dimensions['F'].width = 40
    
    wb.save(out_xlsx)
    print(f"Created human review workbook at: {out_xlsx}")
    print(f"Total rows to verify: {len(rows)}")

if __name__ == "__main__":
    main()
