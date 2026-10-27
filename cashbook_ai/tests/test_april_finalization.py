import json
from pathlib import Path

import openpyxl

from finalize_april import HEAD_CORRECTIONS, apply_head_corrections
from process_month import process_month
from src.models import AccountSet, ProcessedRow, RawExtractedRow, RecordType, TransactionType
from src.validation.head_master import HeadMasterRepository


def test_head_corrections_preserve_raw_values_and_use_canonical_names(tmp_path):
    master = tmp_path / "master.csv"
    master.write_text(
        "account_set,accounting_head_number,extracted_head_name,human_final_name,accounting_head_name,source_type\n"
        "GENERAL,20,Original 20,Canonical Twenty,Canonical Twenty,MANUAL_VERIFIED\n"
        "GENERAL,40,Original 40,Canonical Forty,Canonical Forty,MANUAL_VERIFIED\n",
        encoding="utf-8",
    )
    repository = HeadMasterRepository(master)
    raws = [
        RawExtractedRow(
            source_image="page_002.jpg", source_page="2", source_row=3,
            raw_head_number="80", raw_voucher_number="524",
        ),
        RawExtractedRow(
            source_image="page_002.jpg", source_page="2", source_row=4,
            raw_head_number="80", raw_voucher_number="525",
        ),
    ]
    processed = [
        ProcessedRow(
            source_image="page_002.jpg", source_page="2", source_row=3,
            raw_head_number="80", raw_voucher_number="524",
            accounting_head_number="80", transaction_type=TransactionType.PAYMENT,
            record_type=RecordType.TRANSACTION,
        ),
        ProcessedRow(
            source_image="page_002.jpg", source_page="2", source_row=4,
            raw_head_number="80", raw_voucher_number="525",
            accounting_head_number="80", transaction_type=TransactionType.PAYMENT,
            record_type=RecordType.TRANSACTION,
        ),
    ]
    corrected, records = apply_head_corrections(processed, raws, repository)
    assert [row.accounting_head_number for row in corrected] == ["20", "40"]
    assert [row.accounting_head_name for row in corrected] == ["Canonical Twenty", "Canonical Forty"]
    assert [row.raw_head_number for row in corrected] == ["80", "80"]
    assert [item["original_raw_head_number"] for item in records] == ["80", "80"]


def test_ready_workbook_has_hidden_head_dropdown(tmp_path):
    from src.export.manual_review_exporter import ManualReviewExporter

    row = ProcessedRow(
        source_image="page_001.jpg", source_page="1", source_row=1,
        resolved_date="2026-04-01", raw_voucher_number="1",
        accounting_head_name="Canonical Head", payment_amount=10,
        transaction_type=TransactionType.PAYMENT, record_type=RecordType.TRANSACTION,
    )
    output = tmp_path / "ready.xlsx"
    ManualReviewExporter().export_with_directions(
        [row],
        {("page_001.jpg", "1", 1): TransactionType.PAYMENT},
        output,
        canonical_head_names=["Canonical Head", "Other Head", "Canonical Head"],
    )
    workbook = openpyxl.load_workbook(output)
    assert [sheet.title for sheet in workbook.worksheets if sheet.sheet_state == "visible"] == ["Receipts", "Payments"]
    assert workbook["Receipts"].max_column == 6
    assert workbook["Payments"].max_column == 6
    assert workbook["Payments"].data_validations.dataValidation[0].formula1 == "=CanonicalHeadNames"
    assert "__HeadNames" in workbook.sheetnames
    assert workbook["__HeadNames"].sheet_state == "hidden"


def test_monthly_workflow_reuses_non_april_raw_fixture(tmp_path):
    raw = tmp_path / "raw_extractions.json"
    raw.write_text(
        json.dumps(
            [
                {
                    "source_image": "may_page_001.jpg",
                    "source_page": "1",
                    "source_row": 1,
                    "raw_date": "1.05.2026",
                    "raw_head_number": "20",
                    "raw_voucher_number": "1",
                    "raw_narration": "entry",
                    "english_narration": "entry",
                    "raw_amount": "10",
                    "raw_amount_col1": "10",
                    "raw_amount_col2": None,
                    "detected_amount_column": 1,
                }
            ]
        ),
        encoding="utf-8",
    )
    output = tmp_path / "2026-05_ready_for_entry.xlsx"
    report = tmp_path / "2026-05_reconciliation.json"
    result = process_month("2026-05", raw, output, report)
    assert result["month"] == "2026-05"
    assert result["receipts"] == 1
    assert result["payments"] == 0
    assert output.exists()
