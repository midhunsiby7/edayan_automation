"""Comprehensive test suite for Milestone 6: Unified Local Application, Upload Interface,
Processing Orchestration, Intelligent Review, Draft Export, and Approved Workbook Intake.
"""

import io
import json
import shutil
import tempfile
from pathlib import Path
from PIL import Image
import pytest
from fastapi.testclient import TestClient
import openpyxl

from src.app.main import app, DATA_DIR, registry
from src.orchestration.job_store import JobStore, JobRecord, JobStatus, StageStatus
from src.orchestration.processing_pipeline import ProcessingPipeline
from src.extraction.mock_extractor import MockExtractor
from src.mapping.review_manager import ReviewAction, ReviewManager
from src.validation.approved_workbook_validator import ApprovedWorkbookValidator


@pytest.fixture
def client():
    return TestClient(app)


@pytest.fixture
def test_workspace(tmp_path):
    """Create isolated test workspace for month operations."""
    ws_data = tmp_path / "data"
    ws_data.mkdir(parents=True)
    for sub in ["monthly", "output", "completed", "archive", "jobs"]:
        (ws_data / sub).mkdir(parents=True)
    return ws_data


def create_test_image_bytes(color=(255, 255, 255), size=(100, 100)) -> bytes:
    """Generate in-memory valid JPEG bytes."""
    buf = io.BytesIO()
    img = Image.new("RGB", size, color=color)
    img.save(buf, format="JPEG")
    return buf.getvalue()


# =========================================================================
# Part 1 & 2: Startup, Dashboard & Status
# =========================================================================

def test_app_startup_and_status(client):
    """Verify application health, host binding, canonical registry counts, and hardware diagnostics."""
    resp = client.get("/api/status")
    assert resp.status_code == 200
    data = resp.json()

    assert data["status"] == "ONLINE"
    assert data["host"] == "127.0.0.1:8000"
    assert "active_month" in data
    assert "available_months" in data
    assert data["browser_relay_status"] in ("ONLINE", "STANDBY")
    assert "local_llm_model" in data
    assert data["canonical_registry"]["heads_count"] == 60
    assert data["canonical_registry"]["subheads_count"] == 25


def test_months_listing(client):
    """Verify listing available accounting months with status."""
    resp = client.get("/api/months")
    assert resp.status_code == 200
    data = resp.json()
    assert "months" in data
    month_names = [m["month"] for m in data["months"]]
    assert "2026-05" in month_names or "2026-04" in month_names


# =========================================================================
# Part 3: Month Isolation & Directory Partitioning
# =========================================================================

def test_create_month_workspace_valid_and_invalid(client, monkeypatch, tmp_path):
    """Verify month creation partitions directories and enforces YYYY-MM format."""
    monkeypatch.setattr("src.app.main.DATA_DIR", tmp_path)

    # Valid month
    resp = client.post("/api/months", data={"month": "2026-08"})
    assert resp.status_code == 200
    assert (tmp_path / "monthly" / "2026-08").is_dir()
    assert (tmp_path / "output" / "2026-08").is_dir()

    # Invalid format
    resp_bad = client.post("/api/months", data={"month": "invalid-month"})
    assert resp_bad.status_code == 400
    assert "Invalid month format" in resp_bad.json()["detail"]


# =========================================================================
# Part 4: Image Upload, Corruption Check, and Duplicate Detection
# =========================================================================

def test_image_upload_and_corruption_detection(client, monkeypatch, tmp_path):
    """Verify image upload saves files non-destructively and rejects corrupted files."""
    monkeypatch.setattr("src.app.main.DATA_DIR", tmp_path)
    (tmp_path / "monthly" / "2026-09").mkdir(parents=True)

    valid_img = create_test_image_bytes(color=(100, 150, 200))
    corrupt_bytes = b"NOT_A_VALID_IMAGE_DATA_BYTES"

    # 1. Upload valid image
    resp = client.post(
        "/api/upload",
        data={"month": "2026-09"},
        files=[("files", ("cashbook_pg1.jpg", valid_img, "image/jpeg"))]
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["uploaded_count"] == 1
    assert data["pages"][0]["page_number"] == "001"
    assert data["pages"][0]["original_filename"] == "cashbook_pg1.jpg"
    assert data["pages"][0]["is_duplicate"] is False

    # 2. Upload corrupt image
    resp_corrupt = client.post(
        "/api/upload",
        data={"month": "2026-09"},
        files=[("files", ("corrupt.jpg", corrupt_bytes, "image/jpeg"))]
    )
    assert resp_corrupt.status_code == 400
    assert "Corrupted or invalid image" in resp_corrupt.json()["detail"]


def test_duplicate_image_detection(client, monkeypatch, tmp_path):
    """Verify duplicate uploads with identical hashes are detected and flagged."""
    monkeypatch.setattr("src.app.main.DATA_DIR", tmp_path)
    (tmp_path / "monthly" / "2026-09").mkdir(parents=True)

    img_bytes = create_test_image_bytes(color=(50, 100, 150))

    # First upload
    client.post(
        "/api/upload",
        data={"month": "2026-09"},
        files=[("files", ("page_first.jpg", img_bytes, "image/jpeg"))]
    )

    # Second upload with same content but different filename
    resp2 = client.post(
        "/api/upload",
        data={"month": "2026-09"},
        files=[("files", ("page_duplicate.jpg", img_bytes, "image/jpeg"))]
    )
    assert resp2.status_code == 200
    data2 = resp2.json()
    assert data2["total_pages"] == 2
    assert data2["pages"][1]["is_duplicate"] is True
    assert data2["pages"][1]["page_number"] == "002"


# =========================================================================
# Part 5: Page Reordering & Rotation
# =========================================================================

def test_page_reordering_and_rotation(client, monkeypatch, tmp_path):
    """Verify non-destructive reordering assigns continuous page numbers (001, 002...)."""
    monkeypatch.setattr("src.app.main.DATA_DIR", tmp_path)
    month_dir = tmp_path / "monthly" / "2026-09"
    month_dir.mkdir(parents=True)

    img1 = create_test_image_bytes(color=(10, 20, 30))
    img2 = create_test_image_bytes(color=(40, 50, 60))

    # Upload two pages
    client.post(
        "/api/upload",
        data={"month": "2026-09"},
        files=[
            ("files", ("p1.jpg", img1, "image/jpeg")),
            ("files", ("p2.jpg", img2, "image/jpeg")),
        ]
    )

    pages = client.get("/api/pages?month=2026-09").json()["pages"]
    assert pages[0]["original_filename"] == "p1.jpg"
    assert pages[1]["original_filename"] == "p2.jpg"

    # Swap order: p2 first, then p1 with 90 degree rotation
    pages[0]["rotation_degrees"] = 90
    reordered_payload = [pages[1], pages[0]]

    resp = client.post(
        "/api/pages/reorder",
        json={"month": "2026-09", "pages": reordered_payload}
    )
    assert resp.status_code == 200
    updated_pages = resp.json()["pages"]
    assert updated_pages[0]["original_filename"] == "p2.jpg"
    assert updated_pages[0]["page_number"] == "001"
    assert updated_pages[1]["original_filename"] == "p1.jpg"
    assert updated_pages[1]["page_number"] == "002"
    assert updated_pages[1]["rotation_degrees"] == 90


# =========================================================================
# Part 6: Processing Job Persistence, Cancellation, and Restart Recovery
# =========================================================================

def test_job_store_persistence_and_cancellation(tmp_path):
    """Verify atomic disk writes, stage updates, and cancellation."""
    jobs_dir = tmp_path / "jobs"
    store = JobStore(jobs_dir=jobs_dir)

    job = store.create_job(month="2026-09")
    assert job.status == JobStatus.PENDING
    assert len(job.stages) == 9

    # Verify persisted to disk
    job_file = jobs_dir / f"{job.job_id}.json"
    assert job_file.is_file()

    # Cancel job
    store.request_cancellation(job.job_id)
    reloaded = store.get_job(job.job_id)
    assert reloaded.status == JobStatus.CANCELLED
    assert reloaded.cancel_requested is True


def test_job_store_crash_recovery(tmp_path):
    """Verify interrupted RUNNING jobs are automatically recovered on application startup."""
    jobs_dir = tmp_path / "jobs"
    jobs_dir.mkdir(parents=True)

    stale_record = {
        "job_id": "job_202609_testcrash",
        "month": "2026-09",
        "status": "RUNNING",
        "current_stage": "ocr_extraction",
        "stages": [],
        "created_at": "2026-10-30T10:00:00Z",
        "updated_at": "2026-10-30T10:01:00Z"
    }
    (jobs_dir / "job_202609_testcrash.json").write_text(json.dumps(stale_record), encoding="utf-8")

    # Startup new JobStore instance
    new_store = JobStore(jobs_dir=jobs_dir)
    recovered = new_store.get_job("job_202609_testcrash")
    assert recovered.status == JobStatus.FAILED
    assert "interrupted by application restart" in recovered.error_message


def test_processing_pipeline_execution_with_mock(tmp_path):
    """Verify complete 9-stage pipeline execution and draft workbook generation."""
    jobs_dir = tmp_path / "jobs"
    output_dir = tmp_path / "output" / "2026-09"
    monthly_dir = tmp_path / "monthly" / "2026-09"
    jobs_dir.mkdir(parents=True)
    output_dir.mkdir(parents=True)
    monthly_dir.mkdir(parents=True)

    # Create dummy images
    img1_path = monthly_dir / "p1.jpg"
    img2_path = monthly_dir / "p2.jpg"
    img1_path.write_bytes(create_test_image_bytes(color=(1, 2, 3)))
    img2_path.write_bytes(create_test_image_bytes(color=(4, 5, 6)))

    page_manifest = [
        {"page_number": "001", "filename": "p1.jpg", "path": str(img1_path), "sha256": "abc1"},
        {"page_number": "002", "filename": "p2.jpg", "path": str(img2_path), "sha256": "abc2"},
    ]

    store = JobStore(jobs_dir=jobs_dir)
    job = store.create_job(month="2026-09")

    proc_pipeline = ProcessingPipeline(
        registry=registry,
        job_store=store,
        data_dir=tmp_path
    )

    completed_job = proc_pipeline.run_pipeline(
        month="2026-09",
        job_id=job.job_id,
        page_manifest=page_manifest,
        use_llm=False,
        force_ocr=True,
        custom_extractor=MockExtractor()
    )

    assert completed_job.status == JobStatus.COMPLETED
    assert completed_job.result_summary["total_rows"] > 0
    assert (output_dir / "2026-09_draft_review.xlsx").is_file()
    assert (output_dir / "review_state.json").is_file()


# =========================================================================
# Part 7: Review Intelligence & User Actions
# =========================================================================

def test_review_decisions_and_alias_promotion(client, monkeypatch, tmp_path):
    """Verify Confirm, Reject, Override, and permanent alias promotion."""
    monkeypatch.setattr("src.app.main.DATA_DIR", tmp_path)
    output_dir = tmp_path / "output" / "2026-09"
    output_dir.mkdir(parents=True)

    # Create initial review state
    review_state = {
        "month": "2026-09",
        "updated_at": "2026-10-30T12:00:00Z",
        "total_items": 1,
        "items": [
            {
                "transaction_id": "tx_test_001",
                "source_image": "p1.jpg",
                "source_page": "001",
                "source_row": 2,
                "raw_narration": "നിത്യചിലവുകൾ സംഭാവന",
                "english_narration": "Daily expenses donation",
                "direction": "RECEIPT",
                "amount": 5000.0,
                "mapping_status": "SUGGESTED",
                "canonical_head_id": "GEN_HEAD_MISC_INCOME",
                "canonical_head_name": "Miscellaneous Receipts",
                "confidence_score": 0.88,
                "needs_human_review": True
            }
        ],
        "groups": []
    }
    (output_dir / "review_state.json").write_text(json.dumps(review_state), encoding="utf-8")

    # 1. Fetch review state
    resp = client.get("/api/review?month=2026-09")
    assert resp.status_code == 200
    assert resp.json()["has_data"] is True
    assert resp.json()["stats"]["suggested"] == 1

    # 2. Confirm transaction
    resp_confirm = client.post(
        "/api/review/tx_test_001/decision",
        json={
            "month": "2026-09",
            "action": "CONFIRM",
            "reviewer": "Accountant"
        }
    )
    assert resp_confirm.status_code == 200
    assert resp_confirm.json()["status"] == "VERIFIED"

    # 3. Override transaction with explicit alias promotion
    resp_override = client.post(
        "/api/review/tx_test_001/decision",
        json={
            "month": "2026-09",
            "action": "OVERRIDE",
            "confirmed_head_id": "GEN_HEAD_DONATION_SPECIAL_C",
            "notes": "Verified against parish receipt",
            "promote_to_alias": True,
            "reviewer": "Accountant"
        }
    )
    assert resp_override.status_code == 200
    item = resp_override.json()["item"]
    assert item["canonical_head_id"] == "GEN_HEAD_DONATION_SPECIAL_C"
    assert item["mapping_status"] == "VERIFIED"
    assert item["alias_promoted"] is True


# =========================================================================
# Part 8: Approved Workbook Intake Validation
# =========================================================================

def create_valid_approved_workbook(path: Path):
    """Create a minimal valid Excel workbook with Receipts and Payments sheets."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws_rec = wb.create_sheet("Receipts")
    ws_rec.append(["Date", "Receipt Number", "Amount", "Head Name", "Subhead", "Narration"])
    ws_rec.append(["2026-05-01", "101", 1000.0, "Donation/Offering for Special Purposes(C)", "", "Parishioner donation"])

    ws_pay = wb.create_sheet("Payments")
    ws_pay.append(["Date", "Receipt Number", "Amount", "Head Name", "Subhead", "Narration"])
    ws_pay.append(["2026-05-02", "201", 500.0, "Electricity Charges", "", "Power bill"])

    wb.save(path)


def test_approved_workbook_intake_valid_and_invalid(client, monkeypatch, tmp_path):
    """Verify intake validates schema, heads, computes SHA-256, and emits PreflightSummary."""
    monkeypatch.setattr("src.app.main.DATA_DIR", tmp_path)
    (tmp_path / "output" / "2026-05" / "intake").mkdir(parents=True)

    valid_wb_path = tmp_path / "valid_approved.xlsx"
    create_valid_approved_workbook(valid_wb_path)

    # Valid upload
    with open(valid_wb_path, "rb") as f:
        resp = client.post(
            "/api/upload-approved",
            data={"month": "2026-05"},
            files={"file": ("valid_approved.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        )
    assert resp.status_code == 200
    summary = resp.json()
    assert summary["is_valid"] is True
    assert summary["status"] == "READY_FOR_EXECUTION"
    assert summary["receipts_count"] == 1
    assert summary["receipts_total"] == 1000.0
    assert summary["payments_count"] == 1
    assert summary["payments_total"] == 500.0
    assert summary["net_movement"] == 500.0
    assert len(summary["file_hash_sha256"]) == 64

    # Corrupt / invalid sheet upload
    bad_wb = openpyxl.Workbook()
    bad_path = tmp_path / "bad.xlsx"
    bad_wb.save(bad_path)

    with open(bad_path, "rb") as f:
        resp_bad = client.post(
            "/api/upload-approved",
            data={"month": "2026-05"},
            files={"file": ("bad.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        )
    assert resp_bad.status_code == 200
    bad_summary = resp_bad.json()
    assert bad_summary["is_valid"] is False
    assert bad_summary["status"] == "VALIDATION_FAILED"
    assert len(bad_summary["errors"]) > 0


# =========================================================================
# Part 9: Security & Path Traversal Prevention
# =========================================================================

def test_security_path_traversal_prevention(client):
    """Verify path traversal attempts are blocked."""
    resp = client.get("/api/pages/image/2026-05/..%2f..%2fsecret.txt")
    assert resp.status_code in (404, 400)
