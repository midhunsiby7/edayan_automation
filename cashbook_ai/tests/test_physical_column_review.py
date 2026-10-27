import json
from pathlib import Path

import openpyxl
import pytest

from src.models import ProcessedRow, RecordType, TransactionType
from src.review.physical_column_review import (
    REVIEW_COLUMNS,
    build_reconciliation_report,
    build_review_records,
    export_verified_workbook,
    import_review_workbook,
    load_decisions,
    write_review_workbook,
)
from correct_april import normalize_direction


def make_row(row=1, transaction_type=TransactionType.PAYMENT, record_type=RecordType.TRANSACTION):
    return ProcessedRow(
        source_image="page_001.jpg",
        source_page="1",
        source_row=row,
        resolved_date="2026-04-01",
        raw_voucher_number=str(row),
        raw_amount_col1="10" if transaction_type == TransactionType.RECEIPT else None,
        raw_amount_col2="20" if transaction_type == TransactionType.PAYMENT else None,
        detected_amount_column=1 if transaction_type == TransactionType.RECEIPT else 2,
        receipt_amount=10 if transaction_type == TransactionType.RECEIPT else None,
        payment_amount=20 if transaction_type == TransactionType.PAYMENT else None,
        accounting_head_name="Verified head",
        narration_english="Narration",
        transaction_type=transaction_type,
        record_type=record_type,
    )


def make_source(tmp_path: Path) -> Path:
    source = tmp_path / "raw.json"
    source.write_text('{"immutable": true}', encoding="utf-8")
    return source


def test_review_workbook_has_all_rows_and_dropdown_fields(tmp_path):
    source = make_source(tmp_path)
    rows = [
        make_row(1),
        make_row(2, TransactionType.RECEIPT),
        make_row(3, record_type=RecordType.OPENING_BALANCE, transaction_type=TransactionType.INFORMATIONAL),
    ]
    records = build_review_records(rows, source)
    output = tmp_path / "review.xlsx"
    write_review_workbook(records, output, source_path=source, source_hash=__import__("src.review.physical_column_review", fromlist=["extraction_sha256"]).extraction_sha256(source))
    workbook = openpyxl.load_workbook(output)
    assert workbook.sheetnames == ["Physical_Column_Review", "Review_Metadata"]
    assert [cell.value for cell in workbook["Physical_Column_Review"][1]] == REVIEW_COLUMNS
    assert len(workbook["Physical_Column_Review"].data_validations.dataValidation) == 2


def test_agreement_does_not_approve_and_unreviewed_blocks(tmp_path):
    source = make_source(tmp_path)
    rows = [make_row()]
    records = build_review_records(rows, source)
    decisions = {
        (rows[0].source_image, rows[0].source_page, rows[0].source_row): {
            "review_status": "UNREVIEWED",
            "human_verified_direction": None,
        }
    }
    report = build_reconciliation_report(rows, decisions)
    assert report["finalization_allowed"] is False
    assert report["unreviewed_count"] == 1


def test_human_verified_direction_overrides_model_without_mutating_row(tmp_path):
    source = make_source(tmp_path)
    rows = [make_row(transaction_type=TransactionType.PAYMENT)]
    key = (rows[0].source_image, rows[0].source_page, rows[0].source_row)
    decisions = {
        key: {
            "review_status": "HUMAN_VERIFIED",
            "human_verified_direction": "Receipt",
        }
    }
    output = tmp_path / "verified.xlsx"
    export_verified_workbook(rows, decisions, output)
    workbook = openpyxl.load_workbook(output, read_only=True, data_only=True)
    assert workbook.sheetnames == ["Receipts", "Payments"]
    assert workbook["Receipts"].max_row == 2
    assert workbook["Payments"].max_row == 1
    assert rows[0].transaction_type == TransactionType.PAYMENT


def test_corrected_direction_uses_original_amount_when_model_side_differs(tmp_path):
    output = tmp_path / "verified.xlsx"
    row = make_row(transaction_type=TransactionType.PAYMENT)
    key = (row.source_image, row.source_page, row.source_row)
    export_verified_workbook(
        [row],
        {key: {"review_status": "HUMAN_VERIFIED", "human_verified_direction": "Receipt"}},
        output,
    )
    workbook = openpyxl.load_workbook(output, read_only=True, data_only=True)
    assert workbook["Receipts"].cell(2, 3).value == 20


def test_saved_direction_normalization_does_not_use_semantics():
    assert normalize_direction("Receipts") == "Receipt"
    assert normalize_direction("PAYMENTS") == "Payment"
    assert normalize_direction("electricity expense") is None


def test_model_disagreement_requires_correction_reason(tmp_path):
    source = make_source(tmp_path)
    rows = [make_row(transaction_type=TransactionType.PAYMENT)]
    records = build_review_records(rows, source)
    records[0]["human_verified_direction"] = "Receipt"
    records[0]["review_status"] = "HUMAN_VERIFIED"
    output = tmp_path / "review.xlsx"
    from src.review.physical_column_review import extraction_sha256
    write_review_workbook(records, output, source_path=source, source_hash=extraction_sha256(source))
    with pytest.raises(ValueError, match="correction reason"):
        import_review_workbook(output, source, rows, tmp_path / "sidecar.json")


def test_import_rejects_stale_hash(tmp_path):
    source = make_source(tmp_path)
    rows = [make_row()]
    records = build_review_records(rows, source)
    workbook_path = tmp_path / "review.xlsx"
    write_review_workbook(records, workbook_path, source_path=source, source_hash="correct")
    source.write_text('{"immutable": false}', encoding="utf-8")
    with pytest.raises(ValueError, match="hash"):
        import_review_workbook(workbook_path, source, rows, tmp_path / "sidecar.json")


def test_import_rejects_duplicate_source_identifier(tmp_path):
    source = make_source(tmp_path)
    rows = [make_row(1), make_row(2)]
    records = build_review_records(rows, source)
    workbook_path = tmp_path / "review.xlsx"
    from src.review.physical_column_review import extraction_sha256
    write_review_workbook(records, workbook_path, source_path=source, source_hash=extraction_sha256(source))
    workbook = openpyxl.load_workbook(workbook_path)
    sheet = workbook["Physical_Column_Review"]
    sheet.cell(3, 1).value = sheet.cell(2, 1).value
    sheet.cell(3, 2).value = sheet.cell(2, 2).value
    sheet.cell(3, 3).value = sheet.cell(2, 3).value
    workbook.save(workbook_path)
    with pytest.raises(ValueError, match="Duplicate"):
        import_review_workbook(workbook_path, source, rows, tmp_path / "sidecar.json")
