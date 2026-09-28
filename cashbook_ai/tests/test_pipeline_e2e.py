import openpyxl
from pathlib import Path
from src.export.excel_exporter import ExcelExporter
from src.extraction.mock_extractor import MockExtractor
from src.models import RecordType, TransactionType, AccountSet, AccountingHead
from src.normalization.date_normalizer import DateNormalizer
from src.validation.accounting_rules import AccountingRuleEngine
from src.validation.head_master import HeadMasterRepository


def test_mock_pipeline_e2e(tmp_path: Path):
    """Verify end-to-end extraction, normalization, accounting rules, and Excel export."""
    extractor = MockExtractor()
    date_normalizer = DateNormalizer()
    
    # Setup HeadMaster
    repo_file = tmp_path / "head_master.json"
    repo = HeadMasterRepository(repo_file)
    repo.add_record(AccountingHead(account_set=AccountSet.GENERAL, accounting_head_number="78", accounting_head_name="Mock Head 78"))
    
    rule_engine = AccountingRuleEngine(date_normalizer, head_master_repo=repo)
    exporter = ExcelExporter()

    # Process Page 1
    p1_rows = extractor.extract(Path("cashbook_page1.jpg"))
    processed_p1 = [rule_engine.process_row(r) for r in p1_rows]

    # Process Page 2 (Cross-page date continuity test)
    p2_rows = extractor.extract(Path("cashbook_page2.jpg"))
    processed_p2 = [rule_engine.process_row(r) for r in p2_rows]

    all_processed = processed_p1 + processed_p2
    assert len(all_processed) == 8

    # Verify Page 1 Row 1: Opening balance
    assert all_processed[0].record_type == RecordType.OPENING_BALANCE
    assert all_processed[0].receipt_amount == 25000.00
    assert all_processed[0].voucher_number is None

    # Verify Page 1 Row 2: Head 78 / Voucher 1 / Payment 700
    assert all_processed[1].accounting_head_number == "78"
    assert all_processed[1].accounting_head_name == "Mock Head 78"
    assert all_processed[1].voucher_number == "1"
    assert all_processed[1].payment_amount == 700.00
    assert all_processed[1].transaction_type == TransactionType.PAYMENT

    # Verify Page 1 Row 3: Head 78 / Voucher 2 / Payment 1080 (Inherited Date 2024-04-02)
    assert all_processed[2].accounting_head_number == "78"
    assert all_processed[2].voucher_number == "2"
    assert all_processed[2].payment_amount == 1080.00
    assert all_processed[2].resolved_date == "2024-04-02"

    # Verify Page 2 Row 1: Inherited date from Page 1!
    assert all_processed[6].source_page == "2"
    assert all_processed[6].resolved_date == "2024-04-05"  # last date on Page 1 was 05/04/2024

    # Export to Excel
    out_xlsx = tmp_path / "test_cashbook.xlsx"
    exporter.export(all_processed, out_xlsx)
    assert out_xlsx.exists()

    # Inspect sheets in exported workbook
    wb = openpyxl.load_workbook(out_xlsx)
    sheet_names = wb.sheetnames
    assert "Receipts" in sheet_names
    assert "Payments" in sheet_names
    assert "All_Transactions" in sheet_names
    assert "Extraction_Audit" in sheet_names

    ws_receipts = wb["Receipts"]
    ws_payments = wb["Payments"]
    ws_all = wb["All_Transactions"]
    ws_audit = wb["Extraction_Audit"]

    # Verify Opening Balance is excluded from Receipts and Payments
    receipt_record_types = [ws_receipts.cell(row=r, column=16).value for r in range(2, ws_receipts.max_row + 1) if ws_receipts.cell(row=r, column=16).value]
    assert RecordType.OPENING_BALANCE.value not in receipt_record_types

    payment_record_types = [ws_payments.cell(row=r, column=16).value for r in range(2, ws_payments.max_row + 1) if ws_payments.cell(row=r, column=16).value]
    assert RecordType.OPENING_BALANCE.value not in payment_record_types

    # Verify Opening Balance is included in All_Transactions and Extraction_Audit
    all_record_types = [ws_all.cell(row=r, column=16).value for r in range(2, ws_all.max_row + 1)]
    assert RecordType.OPENING_BALANCE.value in all_record_types

    audit_record_types = [ws_audit.cell(row=r, column=25).value for r in range(2, ws_audit.max_row + 1)]
    assert RecordType.OPENING_BALANCE.value in audit_record_types

    # Check All_Transactions headers (18 columns including Source Image)
    ws_all = wb["All_Transactions"]
    header_vals = [ws_all.cell(row=1, column=c).value for c in range(1, 19)]
    assert "Raw Date" in header_vals
    assert "Resolved Date" in header_vals
    assert "Acct Head No" in header_vals
    assert "Acct Head Name" in header_vals
    assert "Web Head Value" in header_vals
    assert "Web Head Label" in header_vals
    assert "Voucher No" in header_vals
    assert "Raw Narration (Malayalam)" in header_vals
    assert "English Narration" in header_vals
    assert "Receipt Amount" in header_vals
    assert "Payment Amount" in header_vals
    assert "Transaction Type" in header_vals
    assert "Record Type" in header_vals
    assert "Confidence" in header_vals
    assert "Review Required" in header_vals

    # Check Extraction_Audit sheet has granular confidence headers
    ws_audit = wb["Extraction_Audit"]
    audit_headers = [ws_audit.cell(row=1, column=c).value for c in range(1, 29)]
    assert "Date Conf" in audit_headers
    assert "Head Conf" in audit_headers
    assert "Voucher Conf" in audit_headers
    assert "Narration Conf" in audit_headers
    assert "Amount Conf" in audit_headers
    assert "Overall Conf" in audit_headers
    assert "Review Reasons / Warnings" in audit_headers
