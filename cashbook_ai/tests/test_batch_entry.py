"""Tests for Guarded Batch Execution Mode and Manifests."""

import json
from pathlib import Path
import sys
from typing import Any, Mapping

import openpyxl
import pytest

from automation.batch_manifest import (
    APPROVED,
    DRAFT,
    BatchManifest,
    ManifestError,
    ManifestStateError,
    UnauthorizedTransactionError,
    WorkbookTamperedError,
    approve_manifest,
    calculate_file_sha256,
    check_source_completeness,
    generate_batch_manifest,
    load_manifest,
    save_manifest,
    validate_manifest,
)
from automation.edayan_entry import (
    CONFIRMED,
    FAILED,
    PENDING,
    SUBMITTED,
    UNCERTAIN,
    BatchEntryRunner,
    ReconciliationError,
    RelayBrowserEntryAdapter,
    check_live_duplicate,
    load_or_create_journal,
    main,
    reconcile_uncertain_entry,
    transition,
    verify_readback,
    verify_table_row_evidence,
    wait_for_subhead,
)


class MockBrowserEntryAdapter:
    """Mock adapter simulating browser interactions without network/DOM access."""

    def __init__(
        self,
        *,
        readback_override: Mapping[str, Any] | None = None,
        table_text_override: str | None = None,
        save_confirmed: bool = True,
        save_error: str | None = None,
        fail_submit: bool = False,
    ):
        self.readback_override = readback_override
        self.table_text_override = table_text_override
        self.save_confirmed = save_confirmed
        self.save_error = save_error
        self.fail_submit = fail_submit
        self.entered_entries: list[dict[str, Any]] = []
        self.submitted_count = 0
        self.current_entry: dict[str, Any] | None = None

    def enter(self, entry: Mapping[str, Any]) -> Mapping[str, Any]:
        self.current_entry = dict(entry)
        self.entered_entries.append(dict(entry))
        return self.readback(entry)

    def readback(self, entry: Mapping[str, Any] | None = None) -> Mapping[str, Any]:
        target = dict(entry or self.current_entry or {})
        if self.readback_override is not None:
            res = dict(self.readback_override)
            if "table_text" not in res:
                res["table_text"] = self.table_text_override or ""
            return res

        date_val = str(target.get("date", "2026-04-04"))
        if len(date_val) == 10 and date_val[4] == "-" and date_val[7] == "-":
            date_val = f"{date_val[8:10]}-{date_val[5:7]}-{date_val[0:4]}"

        return {
            "date": date_val,
            "receipt_number": str(target.get("receipt_number", "515")),
            "amount": str(target.get("amount", "1000")),
            "mode_cash": True,
            "head": target.get("head_name", "Test Head"),
            "subhead": target.get("subhead", ""),
            "narration": str(target.get("narration", "")),
            "table_text": self.table_text_override or "",
            "dom": {},
        }

    def submit(self) -> Mapping[str, Any]:
        if self.fail_submit:
            raise RuntimeError("Browser submit error")
        self.submitted_count += 1
        return {"clicked": True}

    def observe_save(
        self, expected_entry: Mapping[str, Any] | None = None, timeout_seconds: float = 30
    ) -> Mapping[str, Any]:
        v_num = str((expected_entry or {}).get("receipt_number", "515"))
        if not self.save_confirmed:
            return {
                "confirmed": False,
                "record_id": None,
                "last_error": self.save_error or "Timeout waiting for table evidence",
            }
        return {
            "confirmed": True,
            "record_id": v_num,
            "success_text": f"Saved voucher {v_num}",
            "table_evidence": True,
        }


def make_dummy_workbook(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Receipts"
    ws1.append(["Date", "Receipt Number", "Amount", "Head Name", "Subhead", "Narration"])
    ws1.append(["2026-04-03", "514", 13245, "Head A", "Sub A", "Narr 514"])
    ws1.append(["2026-04-03", "1", 725, "Bank Interest", "", "Narr 1"])
    ws1.append(["2026-04-04", "515", 6405, "Head A", "Sub B", "Narr 515"])
    ws1.append(["2026-04-05", "516", 8000, "Head B", "", "Narr 516"])
    ws2 = wb.create_sheet(title="Payments")
    ws2.append(["Date", "Receipt Number", "Amount", "Head Name", "Subhead", "Narration"])
    ws2.append(["2026-04-01", "1", 700, "Expense Head", "", "Payment 1"])
    wb.save(path)
    return path


def make_sample_report(entries: list[dict[str, Any]]) -> dict[str, Any]:
    return {"mode": "dry-run", "entries": entries}


def test_batch_manifest_generation_and_exclusions(tmp_path: Path):
    wb_file = make_dummy_workbook(tmp_path / "test_wb.xlsx")
    journal_data = {
        "entries": {
            "key-514": {"status": "confirmed", "evidence": {"record_id": "514"}},
        }
    }
    entries = [
        {
            "submission_key": "key-514",
            "direction": "RECEIPT",
            "sheet": "Receipts",
            "row": 2,
            "date": "2026-04-03",
            "receipt_number": "514",
            "amount": 13245,
            "head_name": "Head A",
            "subhead": "Sub A",
            "narration": "Narr 514",
            "status": "READY",
        },
        {
            "submission_key": "key-1",
            "direction": "RECEIPT",
            "sheet": "Receipts",
            "row": 3,
            "date": "2026-04-03",
            "receipt_number": "1",
            "amount": 725,
            "head_name": "Bank Interest",
            "subhead": "",
            "narration": "Narr 1",
            "status": "READY",
        },
        {
            "submission_key": "key-515",
            "direction": "RECEIPT",
            "sheet": "Receipts",
            "row": 4,
            "date": "2026-04-04",
            "receipt_number": "515",
            "amount": 6405,
            "head_name": "Head A",
            "subhead": "Sub B",
            "narration": "Narr 515",
            "status": "READY",
        },
        {
            "submission_key": "key-516",
            "direction": "RECEIPT",
            "sheet": "Receipts",
            "row": 5,
            "date": "2026-04-05",
            "receipt_number": "516",
            "amount": 8000,
            "head_name": "Head B",
            "subhead": "",
            "narration": "Narr 516",
            "status": "READY",
        },
        {
            "submission_key": "key-pay-1",
            "direction": "PAYMENT",
            "sheet": "Payments",
            "row": 2,
            "date": "2026-04-01",
            "receipt_number": "1",
            "amount": 700,
            "head_name": "Expense Head",
            "subhead": "",
            "narration": "Payment 1",
            "status": "READY",
        },
        {
            "submission_key": "key-blocked",
            "direction": "RECEIPT",
            "sheet": "Receipts",
            "row": 6,
            "date": "2026-04-06",
            "receipt_number": "517",
            "amount": 500,
            "head_name": "Unmapped Head",
            "subhead": "",
            "narration": "Blocked",
            "status": "BLOCKED",
            "issues": ["No website mapping"],
        },
    ]

    report = make_sample_report(entries)
    manifest = generate_batch_manifest(
        report, journal_data, wb_file, direction="RECEIPT", target_voucher_start=515, target_voucher_end=516
    )

    assert manifest.status == DRAFT
    assert manifest.batch_summary["transaction_count"] == 2
    assert manifest.batch_summary["total_amount"] == 14405.0  # 6405 + 8000
    assert manifest.authorized_keys == ["key-515", "key-516"]

    # Verify exclusions
    excluded_keys = {e["submission_key"]: e["reason"] for e in manifest.excluded_entries}
    assert "ALREADY_CONFIRMED" in excluded_keys["key-514"]
    assert "VOUCHER_RANGE_EXCLUDED" in excluded_keys["key-1"] or "VOUCHER_DISCONTINUITY" in excluded_keys["key-1"]
    assert "DIRECTION_EXCLUDED" in excluded_keys["key-pay-1"]
    assert "BLOCKED" in excluded_keys["key-blocked"]


def test_batch_manifest_approval_lifecycle(tmp_path: Path):
    wb_file = make_dummy_workbook(tmp_path / "test_wb.xlsx")
    manifest = generate_batch_manifest(
        make_sample_report([]), {}, wb_file, direction="RECEIPT", target_voucher_start=1, target_voucher_end=1
    )
    assert manifest.status == DRAFT
    assert manifest.approval_metadata is None

    # Approve
    approved = approve_manifest(manifest, approved_by="auditor_test")
    assert approved.status == APPROVED
    assert approved.approval_metadata is not None
    assert approved.approval_metadata["approved_by"] == "auditor_test"

    # Cannot approve twice
    with pytest.raises(ManifestStateError, match="already approved"):
        approve_manifest(approved, approved_by="auditor_test")


def test_batch_manifest_rejects_altered_workbook(tmp_path: Path):
    wb_file = make_dummy_workbook(tmp_path / "test_wb.xlsx")
    manifest = generate_batch_manifest(
        make_sample_report([]), {}, wb_file, direction="RECEIPT", target_voucher_start=1, target_voucher_end=1
    )
    validate_manifest(manifest, wb_file)

    # Modify workbook byte
    with wb_file.open("ab") as handle:
        handle.write(b"corruption")

    with pytest.raises(WorkbookTamperedError):
        validate_manifest(manifest, wb_file)


def test_unapproved_draft_manifest_cannot_be_run(tmp_path: Path):
    wb_file = make_dummy_workbook(tmp_path / "test_wb.xlsx")
    manifest = generate_batch_manifest(
        make_sample_report([]), {}, wb_file, direction="RECEIPT", target_voucher_start=1, target_voucher_end=1
    )
    journal_path = tmp_path / "journal.json"
    runner = BatchEntryRunner(manifest, MockBrowserEntryAdapter(), journal_path, wb_file)

    with pytest.raises(ManifestStateError, match="must be APPROVED"):
        runner.run_batch()


def test_whitelist_enforcement_rejects_unauthorized_entry(tmp_path: Path):
    wb_file = make_dummy_workbook(tmp_path / "test_wb.xlsx")
    entry = {
        "submission_key": "key-515",
        "direction": "RECEIPT",
        "sheet": "Receipts",
        "row": 4,
        "date": "2026-04-04",
        "receipt_number": "515",
        "amount": 6405,
        "head_name": "Head A",
        "subhead": "",
        "narration": "Narr 515",
        "status": "READY",
    }
    manifest = generate_batch_manifest(
        make_sample_report([entry]), {}, wb_file, direction="RECEIPT", target_voucher_start=515, target_voucher_end=515
    )
    approved = approve_manifest(manifest, approved_by="user")

    # Infiltrate an unauthorized transaction into approved.authorized_entries (scope expansion attempt)
    approved.authorized_entries.append({
        "submission_key": "unauthorized-injected-key",
        "sheet": "Receipts",
        "row": 10,
        "receipt_number": "999",
        "amount": 100,
    })

    journal_path = tmp_path / "journal.json"
    runner = BatchEntryRunner(approved, MockBrowserEntryAdapter(), journal_path, wb_file)

    with pytest.raises(UnauthorizedTransactionError):
        runner.run_batch()


def test_pre_submit_readback_mismatch_halts_batch(tmp_path: Path):
    wb_file = make_dummy_workbook(tmp_path / "test_wb.xlsx")
    entry = {
        "submission_key": "key-515",
        "direction": "RECEIPT",
        "sheet": "Receipts",
        "row": 4,
        "date": "2026-04-04",
        "receipt_number": "515",
        "amount": 6405,
        "head_name": "Collections as per Order from the Diocese",
        "subhead": "Holy Land",
        "narration": "As per order (Holy land)",
        "status": "READY",
    }
    report = make_sample_report([entry])
    manifest = generate_batch_manifest(
        report, {}, wb_file, direction="RECEIPT", target_voucher_start=515, target_voucher_end=515
    )
    approved = approve_manifest(manifest, approved_by="user")
    journal_path = tmp_path / "journal.json"

    # Adapter simulates readback mismatch (amount typed as 500 instead of 6405)
    adapter = MockBrowserEntryAdapter(
        readback_override={
            "date": "04-04-2026",
            "receipt_number": "515",
            "amount": "500",  # MISMATCH
            "mode_cash": True,
            "head": "Collections as per Order from the Diocese",
            "subhead": "Holy Land",
            "narration": "As per order (Holy land)",
        }
    )

    runner = BatchEntryRunner(approved, adapter, journal_path, wb_file, incident_dir=tmp_path)
    result = runner.run_batch()

    assert result["status"] == "STOPPED"
    assert result["reason"] == "CIRCUIT_BREAKER_TRIPPED"
    assert adapter.submitted_count == 0  # SUBMIT WAS NEVER CLICKED!

    saved_journal = load_or_create_journal(journal_path)
    assert saved_journal["entries"]["key-515"]["status"] == FAILED


def test_duplicate_conflict_in_live_table_halts_batch(tmp_path: Path):
    wb_file = make_dummy_workbook(tmp_path / "test_wb.xlsx")
    entry = {
        "submission_key": "key-515",
        "direction": "RECEIPT",
        "sheet": "Receipts",
        "row": 4,
        "date": "2026-04-04",
        "receipt_number": "515",
        "amount": 6405,
        "head_name": "Head A",
        "subhead": "",
        "narration": "Narr 515",
        "status": "READY",
    }
    manifest = generate_batch_manifest(
        make_sample_report([entry]), {}, wb_file, direction="RECEIPT", target_voucher_start=515, target_voucher_end=515
    )
    approved = approve_manifest(manifest, approved_by="user")
    journal_path = tmp_path / "journal.json"

    # Conflicting line in live table (voucher 515 exists with DIFFERENT amount 9999)
    table_conflict = "1 | 04-04-2026 | Head A | 515 | Cash | Conflict | ₹9,999.00"
    adapter = MockBrowserEntryAdapter(table_text_override=table_conflict)

    runner = BatchEntryRunner(approved, adapter, journal_path, wb_file, incident_dir=tmp_path)
    result = runner.run_batch()

    assert result["status"] == "STOPPED"
    assert adapter.submitted_count == 0  # Did not submit into conflict!

    saved_journal = load_or_create_journal(journal_path)
    assert saved_journal["entries"]["key-515"]["status"] == FAILED


def test_save_timeout_marks_uncertain_and_halts_batch(tmp_path: Path):
    wb_file = make_dummy_workbook(tmp_path / "test_wb.xlsx")
    entry1 = {
        "submission_key": "key-515",
        "direction": "RECEIPT",
        "sheet": "Receipts",
        "row": 4,
        "date": "2026-04-04",
        "receipt_number": "515",
        "amount": 6405,
        "head_name": "Head A",
        "subhead": "",
        "narration": "Narr 515",
        "status": "READY",
    }
    entry2 = {
        "submission_key": "key-516",
        "direction": "RECEIPT",
        "sheet": "Receipts",
        "row": 5,
        "date": "2026-04-05",
        "receipt_number": "516",
        "amount": 8000,
        "head_name": "Head B",
        "subhead": "",
        "narration": "Narr 516",
        "status": "READY",
    }
    manifest = generate_batch_manifest(
        make_sample_report([entry1, entry2]), {}, wb_file, direction="RECEIPT", target_voucher_start=515, target_voucher_end=516
    )
    approved = approve_manifest(manifest, approved_by="user")
    journal_path = tmp_path / "journal.json"

    # Save observation fails/times out
    adapter = MockBrowserEntryAdapter(save_confirmed=False, save_error="Timeout waiting for row")

    runner = BatchEntryRunner(approved, adapter, journal_path, wb_file, incident_dir=tmp_path)
    result = runner.run_batch()

    assert result["status"] == "STOPPED"
    assert result["reason"] == "CIRCUIT_BREAKER_TRIPPED"
    assert adapter.submitted_count == 1  # Only submitted first

    saved_journal = load_or_create_journal(journal_path)
    # First transaction marked UNCERTAIN
    assert saved_journal["entries"]["key-515"]["status"] == UNCERTAIN
    # Second transaction was NOT touched
    assert "key-516" not in saved_journal["entries"]


def test_successful_sequential_batch_execution(tmp_path: Path):
    wb_file = make_dummy_workbook(tmp_path / "test_wb.xlsx")
    entries = [
        {
            "submission_key": "key-515",
            "direction": "RECEIPT",
            "sheet": "Receipts",
            "row": 4,
            "date": "2026-04-04",
            "receipt_number": "515",
            "amount": 6405,
            "head_name": "Head A",
            "subhead": "Sub B",
            "narration": "Narr 515",
            "status": "READY",
        },
        {
            "submission_key": "key-516",
            "direction": "RECEIPT",
            "sheet": "Receipts",
            "row": 5,
            "date": "2026-04-05",
            "receipt_number": "516",
            "amount": 8000,
            "head_name": "Head B",
            "subhead": "",
            "narration": "Narr 516",
            "status": "READY",
        },
    ]
    manifest = generate_batch_manifest(
        make_sample_report(entries), {}, wb_file, direction="RECEIPT", target_voucher_start=515, target_voucher_end=516
    )
    approved = approve_manifest(manifest, approved_by="user")
    journal_path = tmp_path / "journal.json"

    adapter = MockBrowserEntryAdapter(save_confirmed=True)
    runner = BatchEntryRunner(approved, adapter, journal_path, wb_file, incident_dir=tmp_path)
    result = runner.run_batch()

    assert result["status"] == "COMPLETED"
    assert result["confirmed_count"] == 2
    assert adapter.submitted_count == 2

    saved_journal = load_or_create_journal(journal_path)
    assert saved_journal["entries"]["key-515"]["status"] == CONFIRMED
    assert saved_journal["entries"]["key-516"]["status"] == CONFIRMED


def test_exact_duplicate_already_in_table_marks_confirmed_without_resubmit(tmp_path: Path):
    wb_file = make_dummy_workbook(tmp_path / "test_wb.xlsx")
    entry = {
        "submission_key": "key-515",
        "direction": "RECEIPT",
        "sheet": "Receipts",
        "row": 4,
        "date": "2026-04-04",
        "receipt_number": "515",
        "amount": 6405,
        "head_name": "Head A",
        "subhead": "Sub B",
        "narration": "Narr 515",
        "status": "READY",
    }
    manifest = generate_batch_manifest(
        make_sample_report([entry]), {}, wb_file, direction="RECEIPT", target_voucher_start=515, target_voucher_end=515
    )
    approved = approve_manifest(manifest, approved_by="user")
    journal_path = tmp_path / "journal.json"

    # Exact matching line already visible in live table
    matching_table_text = "1 | 04-04-2026 | Head A | 515 | Cash | Narr 515 | ₹6,405.00"
    adapter = MockBrowserEntryAdapter(table_text_override=matching_table_text)

    runner = BatchEntryRunner(approved, adapter, journal_path, wb_file, incident_dir=tmp_path)
    result = runner.run_batch()

    assert result["status"] == "COMPLETED"
    assert result["confirmed_count"] == 1
    assert adapter.submitted_count == 0  # NOT re-submitted! Form submit was skipped safely!

    saved_journal = load_or_create_journal(journal_path)
    assert saved_journal["entries"]["key-515"]["status"] == CONFIRMED


def test_verify_readback_field_validations():
    entry = {
        "date": "2026-04-04",
        "receipt_number": "515",
        "amount": 6405,
        "head_name": "Collections as per Order from the Diocese",
        "subhead": "Holy Land",
        "narration": "As per order (Holy land)",
        "head_resolution": {"website_label": "Collections as per Order from the Diocese ( Credit )"},
    }
    # 1. Exact match
    valid_rb = {
        "date": "04-04-2026",
        "receipt_number": "515",
        "amount": "6405",
        "mode_cash": True,
        "head": "Collections as per Order from the Diocese ( Credit )",
        "subhead": "Holy Land",
        "narration": "As per order (Holy land)",
    }
    is_valid, issues = verify_readback(valid_rb, entry)
    assert is_valid
    assert len(issues) == 0

    # 2. Date mismatch
    bad_date = dict(valid_rb, date="05-04-2026")
    is_valid, issues = verify_readback(bad_date, entry)
    assert not is_valid
    assert any("Date mismatch" in i for i in issues)

    # 3. Voucher mismatch
    bad_v = dict(valid_rb, receipt_number="516")
    is_valid, issues = verify_readback(bad_v, entry)
    assert not is_valid
    assert any("Voucher number mismatch" in i for i in issues)

    # 4. Mode cash unchecked
    bad_mode = dict(valid_rb, mode_cash=False)
    is_valid, issues = verify_readback(bad_mode, entry)
    assert not is_valid
    assert any("Mode cash is not selected" in i for i in issues)


def test_manifest_save_and_load_roundtrip(tmp_path: Path):
    wb_file = make_dummy_workbook(tmp_path / "test_wb.xlsx")
    manifest = generate_batch_manifest(
        make_sample_report([]), {}, wb_file, direction="RECEIPT", target_voucher_start=1, target_voucher_end=1
    )
    manifest_path = tmp_path / "manifest.json"
    save_manifest(manifest, manifest_path)
    loaded = load_manifest(manifest_path)

    assert loaded.manifest_id == manifest.manifest_id
    assert loaded.workbook_sha256 == manifest.workbook_sha256
    assert loaded.status == manifest.status
    assert loaded.batch_summary == manifest.batch_summary


# ==============================================================================
# Objective 1 Tests: Dependent Subhead Settle & Race Prevention
# ==============================================================================

def test_subhead_with_subheads_settles_and_selects():
    """Verify that a head with subheads waits for options to settle and returns state='options'."""
    states = iter([
        {"html": '<select name="subhead" id="subhead"></select>'},
        {"html": '<select name="subhead" id="subhead"><option value="454">Holy Land</option></select>'},
    ])
    result = wait_for_subhead(lambda: next(states), timeout_seconds=2, interval_seconds=0.01)
    assert result["state"] == "options"
    assert "Holy Land" in result["html"]


def test_subhead_no_heads_exist_settles():
    """Verify that a head reporting 'No Heads Exist' settles cleanly to state='no_options'."""
    states = iter([
        {"html": '<select name="subhead" id="subhead"></select>'},
        {"html": '<select name="subhead" id="subhead"><option value="0">No Heads Exist</option></select>'},
    ])
    result = wait_for_subhead(lambda: next(states), timeout_seconds=2, interval_seconds=0.01)
    assert result["state"] == "no_options"
    assert "No Heads Exist" in result["html"]


def test_subhead_delayed_response_settles_with_stability_checks():
    """Verify that a delayed response requiring multiple stable reads settles properly."""
    states = iter([
        {"html": '<select name="subhead"></select>'},
        {"html": '<select name="subhead"></select>'},
        {"html": '<select name="subhead"><option value="454">Holy Land</option></select>'},
        {"html": '<select name="subhead"><option value="454">Holy Land</option></select>'},
    ])
    result = wait_for_subhead(lambda: next(states), timeout_seconds=2, interval_seconds=0.01, stability_checks=2)
    assert result["state"] == "options"


def test_subhead_never_settles_times_out():
    """Verify that an indeterminate subhead control times out safely without proceeding."""
    with pytest.raises(TimeoutError, match="Dependent subhead control did not settle"):
        wait_for_subhead(lambda: {"html": '<select name="subhead"></select>'}, timeout_seconds=0.05, interval_seconds=0.01)


# ==============================================================================
# Objective 2 Tests: Resilient Save Observation
# ==============================================================================

class MockObservableClient:
    """Mock relay client for testing observe_save under navigation and disconnects."""

    def __init__(self, responses: list[Any]):
        self.responses = iter(responses)
        self.command_history: list[tuple[str, dict[str, Any]]] = []

    def get_dom(self, selector: str | None = None) -> dict[str, Any]:
        try:
            resp = next(self.responses)
        except StopIteration:
            return {"url": "https://account.edayan.org/transactions", "title": "EDAYAN", "html": ""}
        if isinstance(resp, Exception):
            raise resp
        return resp

    def command(self, cmd: str, params: dict[str, Any] | None = None) -> dict[str, Any]:
        self.command_history.append((cmd, params or {}))
        return {"data": {"value": "516"}}


def test_observe_save_5_point_table_verification():
    """Verify that verify_table_row_evidence enforces voucher, date, amount, head, and narration."""
    table_text = (
        "1 | 04-04-2026 | Collections as per Order from the DioceseHoly Land | 515 | 💵 | As per order (Holy land) | ₹6,405.00 | 0.00\n"
        "2 | 03-04-2026 | Collections as per Order from the DioceseSeminarian's Aid Fund | 514 | 💵 | As per order (Seminarians aid fund) | ₹13,245.00 | 0.00"
    )
    entry_match = {
        "receipt_number": "515",
        "date": "2026-04-04",
        "amount": 6405,
        "head_name": "Collections as per Order from the Diocese",
        "narration": "As per order (Holy land)",
    }
    is_valid, err, details = verify_table_row_evidence(table_text, entry_match)
    assert is_valid
    assert err is None
    assert details["voucher"] == "515"

    # Amount mismatch
    entry_bad_amt = dict(entry_match, amount=9999)
    is_valid, err, _ = verify_table_row_evidence(table_text, entry_bad_amt)
    assert not is_valid
    assert "amount expected" in err

    # Date mismatch
    entry_bad_date = dict(entry_match, date="2026-04-05")
    is_valid, err, _ = verify_table_row_evidence(table_text, entry_bad_date)
    assert not is_valid
    assert "date expected" in err


def test_observe_save_handles_navigation_disconnection_and_recovers():
    """Verify that observe_save recovers from transient disconnects during full page navigation."""
    entry = {
        "receipt_number": "515",
        "date": "2026-04-04",
        "amount": 6405,
        "head_name": "Collections as per Order from the Diocese",
        "narration": "As per order (Holy land)",
    }
    table_success = "1 | 04-04-2026 | Collections as per Order from the DioceseHoly Land | 515 | 💵 | As per order (Holy land) | ₹6,405.00 | 0.00"

    client = MockObservableClient([
        RuntimeError("Could not establish connection. Receiving end does not exist."),
        RuntimeError("WebSocket connection closed"),
        {"url": "https://account.edayan.org/transactions", "title": "EDAYAN", "html": ""},
        {"text": table_success},
    ])
    adapter = RelayBrowserEntryAdapter(client)
    res = adapter.observe_save(entry, timeout_seconds=5)

    assert res["confirmed"] is True
    assert res["record_id"] == "515"
    assert res["reconnections"] == 2
    assert res["outcome_classification"] == "CONFIRMED_PERSISTENCE"


def test_observe_save_detects_server_rejection():
    """Verify that an alert-danger is detected as CONFIRMED_SERVER_REJECTION."""
    entry = {"receipt_number": "515", "date": "2026-04-04", "amount": 6405}
    html_danger = '<div class="alert alert-danger">Error: Duplicate voucher number detected</div>'

    client = MockObservableClient([
        {"url": "https://account.edayan.org/transactions", "title": "EDAYAN", "html": html_danger},
    ])
    adapter = RelayBrowserEntryAdapter(client)
    res = adapter.observe_save(entry, timeout_seconds=2)

    assert res["confirmed"] is False
    assert res["outcome_classification"] == "CONFIRMED_SERVER_REJECTION"
    assert "Duplicate voucher number" in res["error_text"]


def test_observe_save_detects_contradictory_evidence():
    """Verify that a conflicting record with wrong amount triggers CONTRADICTORY_EVIDENCE."""
    entry = {
        "receipt_number": "515",
        "date": "2026-04-04",
        "amount": 6405,
        "head_name": "Head A",
        "narration": "Narr 515",
    }
    table_conflict = "1 | 04-04-2026 | Head A | 515 | 💵 | Narr 515 | ₹9,999.00 | 0.00"

    client = MockObservableClient([
        {"url": "https://account.edayan.org/transactions", "title": "EDAYAN", "html": ""},
        {"text": table_conflict},
    ])
    adapter = RelayBrowserEntryAdapter(client)
    res = adapter.observe_save(entry, timeout_seconds=1.5)

    assert res["confirmed"] is False
    assert res["outcome_classification"] == "CONTRADICTORY_EVIDENCE"


def test_observe_save_success_alert_plus_correct_five_field_table_evidence():
    """Verify that an alert-success COMBINED with matching 5-point table evidence confirms save."""
    entry = {
        "receipt_number": "515",
        "date": "2026-04-04",
        "amount": 6405,
        "head_name": "Collections as per Order from the Diocese",
        "narration": "As per order (Holy land)",
    }
    html_success = '<div class="alert alert-success">Transaction saved successfully</div>'
    table_success = "1 | 04-04-2026 | Collections as per Order from the DioceseHoly Land | 515 | 💵 | As per order (Holy land) | ₹6,405.00 | 0.00"

    client = MockObservableClient([
        {"url": "https://account.edayan.org/transactions", "title": "EDAYAN", "html": html_success},
        {"text": table_success},
    ])
    adapter = RelayBrowserEntryAdapter(client)
    res = adapter.observe_save(entry, timeout_seconds=2)

    assert res["confirmed"] is True
    assert res["record_id"] == "515"
    assert res["table_evidence"] is True
    assert res["success_text"] == "Transaction saved successfully"
    assert res["outcome_classification"] == "CONFIRMED_PERSISTENCE"


def test_observe_save_success_alert_with_no_matching_table_row_times_out_uncertain():
    """Verify that a success alert alone WITHOUT matching table evidence does NOT confirm and times out to uncertainty."""
    entry = {
        "receipt_number": "516",
        "date": "2026-04-05",
        "amount": 8000,
        "head_name": "Offerings for Daily Expense",
        "narration": "Donation for Holy Mass",
    }
    html_success = '<div class="alert alert-success">Receipt saved successfully</div>'
    table_text = "1 | 04-04-2026 | Collections as per Order | 515 | 💵 | As per order | ₹6,405.00 | 0.00"

    client = MockObservableClient([
        {"url": "https://account.edayan.org/transactions", "title": "EDAYAN", "html": html_success},
        {"text": table_text},
    ])
    adapter = RelayBrowserEntryAdapter(client)
    res = adapter.observe_save(entry, timeout_seconds=1.5)

    assert res["confirmed"] is False
    assert res["table_evidence"] is False
    assert res["success_text"] == "Receipt saved successfully"
    assert res["outcome_classification"] == "UNRESOLVED_UNCERTAIN"


def test_observe_save_success_alert_plus_conflicting_row_halts_with_contradictory_evidence():
    """Verify that a success alert does NOT mask a contradictory table row for the expected voucher."""
    entry = {
        "receipt_number": "516",
        "date": "2026-04-05",
        "amount": 8000,
        "head_name": "Offerings for Daily Expense",
        "narration": "Donation for Holy Mass",
    }
    html_success = '<div class="alert alert-success">Receipt saved successfully</div>'
    table_conflict = "1 | 05-04-2026 | Offerings for Daily Expense | 516 | 💵 | Donation for Holy Mass | ₹9,999.00 | 0.00"

    client = MockObservableClient([
        {"url": "https://account.edayan.org/transactions", "title": "EDAYAN", "html": html_success},
        {"text": table_conflict},
    ])
    adapter = RelayBrowserEntryAdapter(client)
    res = adapter.observe_save(entry, timeout_seconds=1.5)

    assert res["confirmed"] is False
    assert res["outcome_classification"] == "CONTRADICTORY_EVIDENCE"
    assert "amount expected '8,000.00'" in res["contradictory_evidence"]


def test_observe_save_contradictory_evidence_remains_unconfirmed_in_final_return_path():
    """Verify that contradictory evidence cannot be overwritten by any subsequent branch or reset."""
    entry = {
        "receipt_number": "516",
        "date": "2026-04-05",
        "amount": 8000,
        "head_name": "Offerings for Daily Expense",
        "narration": "Donation for Holy Mass",
    }
    table_conflict = "1 | 05-04-2026 | Offerings for Daily Expense | 516 | 💵 | Donation for Holy Mass | ₹1,000.00 | 0.00"

    client = MockObservableClient([
        {"url": "https://account.edayan.org/transactions", "title": "EDAYAN", "html": ""},
        {"text": table_conflict},
    ])
    adapter = RelayBrowserEntryAdapter(client)
    res = adapter.observe_save(entry, timeout_seconds=1.5)

    assert res["confirmed"] is False
    assert res["outcome_classification"] == "CONTRADICTORY_EVIDENCE"
    assert res["record_id"] is None
    assert res["table_evidence"] is False


def test_observe_save_no_duplicate_submission_during_observation_or_recovery():
    """Verify that observe_save only executes read-only inspections and never issues submit commands."""
    entry = {
        "receipt_number": "515",
        "date": "2026-04-04",
        "amount": 6405,
        "head_name": "Collections as per Order from the Diocese",
        "narration": "As per order (Holy land)",
    }
    table_success = "1 | 04-04-2026 | Collections as per Order from the DioceseHoly Land | 515 | 💵 | As per order (Holy land) | ₹6,405.00 | 0.00"

    client = MockObservableClient([
        RuntimeError("Could not establish connection. Receiving end does not exist."),
        RuntimeError("WebSocket closed unexpectedly"),
        {"url": "https://account.edayan.org/transactions", "title": "EDAYAN", "html": ""},
        {"text": table_success},
    ])
    adapter = RelayBrowserEntryAdapter(client)
    res = adapter.observe_save(entry, timeout_seconds=5)

    assert res["confirmed"] is True
    # Verify no click commands (like submit) were issued
    for cmd, params in client.command_history:
        assert cmd != "click", f"Unexpected click command issued during observation: {params}"
        if cmd == "type":
            assert params.get("text") == "", f"Unexpected non-empty type command: {params}"


# ==============================================================================
# Objective 3 Tests: Audited Reconciliation for Uncertain Entries
# ==============================================================================

def test_reconciliation_valid_permit_retry(tmp_path: Path):
    """Verify valid reconciliation transitions UNCERTAIN -> PENDING with full audit trail."""
    journal_path = tmp_path / "journal.json"
    incident_dir = tmp_path / "incidents"
    reconciliation_dir = tmp_path / "reconciliations"
    incident_dir.mkdir()

    sub_key = "key-516"
    inc_id = "incident-1793134120"

    # 1. Setup incident file on disk
    incident_file = incident_dir / f"batch_{inc_id}.json"
    incident_file.write_text(json.dumps({
        "incident_id": inc_id,
        "submission_key": sub_key,
        "voucher_number": "516",
        "amount": 8000,
        "failure_stage": "SAVE_UNCONFIRMED",
        "error_message": "Timeout waiting for table evidence",
        "timestamp": "2026-10-27T20:48:40Z",
    }))

    # 2. Setup journal with UNCERTAIN status
    journal_data = {
        "mode": "live-entry",
        "entries": {
            sub_key: {
                "status": "uncertain",
                "direction": "RECEIPT",
                "sheet": "Receipts",
                "row": 5,
                "voucher_number": "516",
                "head_name": "Head B",
                "error": "Timeout waiting for table evidence",
                "updated_at": "2026-10-27T20:48:40Z",
            }
        }
    }
    journal_path.write_text(json.dumps(journal_data))

    # 3. Perform audited reconciliation
    evidence = {
        "server_table_voucher_absent": True,
        "server_next_receipt_no": "516",
        "cash_closing_balance": "₹53,032.60",
        "daybook_entry_count": 2,
    }
    res = reconcile_uncertain_entry(
        journal_path,
        sub_key,
        incident_id=inc_id,
        authorizer="human_supervisor",
        decision="PERMIT_RETRY",
        evidence=evidence,
        incident_dir=incident_dir,
        reconciliation_dir=reconciliation_dir,
        notes="Absence confirmed on live server",
    )

    assert res["status"] == "RECONCILED"
    assert res["target_status"] == "pending"
    assert Path(res["artifact_path"]).is_file()

    # Verify journal entry was updated
    j_after = json.loads(journal_path.read_text())
    entry_after = j_after["entries"][sub_key]
    assert entry_after["status"] == "pending"
    assert entry_after["prior_status"] == "uncertain"
    assert "error" not in entry_after
    assert len(entry_after["uncertainty_history"]) == 1
    assert entry_after["uncertainty_history"][0]["incident_id"] == inc_id
    assert entry_after["reconciliation"]["authorizer"] == "human_supervisor"
    assert entry_after["reconciliation"]["decision"] == "PERMIT_RETRY"


def test_reconciliation_fails_if_incident_missing(tmp_path: Path):
    journal_path = tmp_path / "journal.json"
    journal_path.write_text(json.dumps({"mode": "live-entry", "entries": {"k1": {"status": "uncertain"}}}))

    with pytest.raises(ReconciliationError, match="Incident 'inc-999' not found"):
        reconcile_uncertain_entry(
            journal_path,
            "k1",
            incident_id="inc-999",
            authorizer="user",
            decision="PERMIT_RETRY",
            evidence={"ok": True},
            incident_dir=tmp_path / "empty_dir",
        )


def test_reconciliation_fails_if_key_or_voucher_mismatch(tmp_path: Path):
    incident_dir = tmp_path / "inc"
    incident_dir.mkdir()
    inc_file = incident_dir / "batch_incident_123.json"
    inc_file.write_text(json.dumps({
        "incident_id": "inc-123",
        "submission_key": "actual-key",
        "voucher_number": "516",
    }))

    journal_path = tmp_path / "journal.json"
    journal_path.write_text(json.dumps({"mode": "live-entry", "entries": {"wrong-key": {"status": "uncertain"}}}))

    with pytest.raises(ReconciliationError, match="does not match requested"):
        reconcile_uncertain_entry(
            journal_path,
            "wrong-key",
            incident_id="inc-123",
            authorizer="user",
            decision="PERMIT_RETRY",
            evidence={"ok": True},
            incident_dir=incident_dir,
        )


def test_reconciliation_fails_if_not_uncertain(tmp_path: Path):
    incident_dir = tmp_path / "inc"
    incident_dir.mkdir()
    (incident_dir / "batch_incident_1.json").write_text(json.dumps({
        "incident_id": "inc-1",
        "submission_key": "k1",
        "voucher_number": "516",
    }))

    journal_path = tmp_path / "journal.json"
    journal_path.write_text(json.dumps({"mode": "live-entry", "entries": {"k1": {"status": "confirmed"}}}))

    with pytest.raises(ReconciliationError, match="status must be 'uncertain'"):
        reconcile_uncertain_entry(
            journal_path,
            "k1",
            incident_id="inc-1",
            authorizer="user",
            decision="PERMIT_RETRY",
            evidence={"ok": True},
            incident_dir=incident_dir,
        )


def test_reconciliation_fails_if_missing_authorizer_or_evidence(tmp_path: Path):
    journal_path = tmp_path / "journal.json"
    incident_dir = tmp_path / "inc"
    incident_dir.mkdir()
    (incident_dir / "batch_incident_1.json").write_text(json.dumps({
        "incident_id": "inc-1",
        "submission_key": "k1",
        "voucher_number": "516",
    }))
    journal_path.write_text(json.dumps({"mode": "live-entry", "entries": {"k1": {"status": "uncertain"}}}))

    # Empty authorizer
    with pytest.raises(ReconciliationError, match="Explicit human authorizer"):
        reconcile_uncertain_entry(
            journal_path, "k1", incident_id="inc-1", authorizer="  ", decision="PERMIT_RETRY", evidence={"a": 1}, incident_dir=incident_dir
        )

    # Empty evidence
    with pytest.raises(ReconciliationError, match="Non-empty evidence"):
        reconcile_uncertain_entry(
            journal_path, "k1", incident_id="inc-1", authorizer="user", decision="PERMIT_RETRY", evidence={}, incident_dir=incident_dir
        )


def test_direct_transition_from_uncertain_to_pending_is_rejected(tmp_path: Path):
    """Enforce that transition() rejects bypassing audited reconciliation."""
    journal_path = tmp_path / "journal.json"
    journal_path.write_text(json.dumps({"mode": "live-entry", "entries": {"k1": {"status": "uncertain"}}}))

    with pytest.raises(ValueError, match="Cannot transition terminal entry"):
        transition(journal_path, "k1", PENDING)


def _setup_reconciliation_fixtures(tmp_path: Path):
    inc_dir = tmp_path / "data" / "output"
    inc_dir.mkdir(parents=True)
    inc_id = "incident-1793134120"
    sub_key = "key-516"

    inc_file = inc_dir / f"batch_{inc_id}.json"
    inc_file.write_text(json.dumps({
        "incident_id": inc_id,
        "submission_key": sub_key,
        "voucher_number": "516",
        "amount": 8000,
        "failure_stage": "SAVE_UNCONFIRMED",
        "error_message": "Timeout waiting for table evidence",
        "timestamp": "2026-10-27T20:48:40Z",
    }))

    journal_file = inc_dir / "journal.json"
    journal_file.write_text(json.dumps({
        "mode": "live-entry",
        "entries": {
            sub_key: {
                "status": "uncertain",
                "direction": "RECEIPT",
                "sheet": "Receipts",
                "row": 5,
                "voucher_number": "516",
                "head_name": "Offerings for Daily Expense",
                "error": "Timeout waiting for table evidence",
                "updated_at": "2026-10-27T20:48:40Z",
            }
        }
    }))

    evidence_file = inc_dir / "evidence.json"
    evidence_file.write_text(json.dumps({"verified_absent": True}))

    return inc_id, sub_key, journal_file, evidence_file, inc_dir


def test_cli_reconciliation_missing_authorizer_rejected(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture):
    inc_id, sub_key, journal_file, evidence_file, inc_dir = _setup_reconciliation_fixtures(tmp_path)
    initial_journal_bytes = journal_file.read_bytes()

    args = [
        "edayan_entry.py",
        "--reconcile-uncertain",
        "--journal", str(journal_file),
        "--submission-key", sub_key,
        "--incident-id", inc_id,
        "--evidence-json", str(evidence_file),
    ]
    monkeypatch.setattr(sys, "argv", args)
    ret = main()

    assert ret != 0
    captured = capsys.readouterr()
    assert "Explicit human authorizer name (--authorizer) is required" in captured.out or "Explicit human authorizer name (--authorizer) is required" in captured.err
    assert journal_file.read_bytes() == initial_journal_bytes
    assert list(inc_dir.glob("reconciliation-*.json")) == []


def test_cli_reconciliation_blank_authorizer_rejected(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture):
    inc_id, sub_key, journal_file, evidence_file, inc_dir = _setup_reconciliation_fixtures(tmp_path)
    initial_journal_bytes = journal_file.read_bytes()

    args = [
        "edayan_entry.py",
        "--reconcile-uncertain",
        "--journal", str(journal_file),
        "--submission-key", sub_key,
        "--incident-id", inc_id,
        "--authorizer", "",
        "--evidence-json", str(evidence_file),
    ]
    monkeypatch.setattr(sys, "argv", args)
    ret = main()

    assert ret != 0
    captured = capsys.readouterr()
    assert "Explicit human authorizer name (--authorizer) is required" in captured.out or "Explicit human authorizer name (--authorizer) is required" in captured.err
    assert journal_file.read_bytes() == initial_journal_bytes
    assert list(inc_dir.glob("reconciliation-*.json")) == []


def test_cli_reconciliation_whitespace_authorizer_rejected(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture):
    inc_id, sub_key, journal_file, evidence_file, inc_dir = _setup_reconciliation_fixtures(tmp_path)
    initial_journal_bytes = journal_file.read_bytes()

    args = [
        "edayan_entry.py",
        "--reconcile-uncertain",
        "--journal", str(journal_file),
        "--submission-key", sub_key,
        "--incident-id", inc_id,
        "--authorizer", "   \t  ",
        "--evidence-json", str(evidence_file),
    ]
    monkeypatch.setattr(sys, "argv", args)
    ret = main()

    assert ret != 0
    captured = capsys.readouterr()
    assert "Explicit human authorizer name (--authorizer) is required" in captured.out or "Explicit human authorizer name (--authorizer) is required" in captured.err
    assert journal_file.read_bytes() == initial_journal_bytes
    assert list(inc_dir.glob("reconciliation-*.json")) == []


def test_cli_reconciliation_valid_authorizer_succeeds(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture):
    inc_id, sub_key, journal_file, evidence_file, inc_dir = _setup_reconciliation_fixtures(tmp_path)

    args = [
        "edayan_entry.py",
        "--reconcile-uncertain",
        "--journal", str(journal_file),
        "--submission-key", sub_key,
        "--incident-id", inc_id,
        "--authorizer", "supervisor_jane",
        "--evidence-json", str(evidence_file),
    ]
    monkeypatch.setattr(sys, "argv", args)
    monkeypatch.chdir(tmp_path)
    ret = main()

    assert ret == 0
    captured = capsys.readouterr()
    result = json.loads(captured.out)
    assert result["status"] == "RECONCILED"
    assert result["reconciliation_record"]["authorizer"] == "supervisor_jane"

    journal_after = json.loads(journal_file.read_text())
    assert journal_after["entries"][sub_key]["status"] == "pending"
    assert journal_after["entries"][sub_key]["reconciliation"]["authorizer"] == "supervisor_jane"
    assert len(list(inc_dir.glob("reconciliation-*.json"))) == 1


def test_extract_table_text_and_verify_html_table():
    from automation.edayan_entry import extract_table_text, verify_table_row_evidence

    html = """
    <table id="myTable">
      <tr>
        <td>1</td>
        <td>05-04-2026</td>
        <td>Offerings for Daily Expense(C)</td>
        <td>516</td>
        <td>💵</td>
        <td>Donation for Holy Mass</td>
        <td>₹8,000.00</td>
        <td>0.00</td>
      </tr>
    </table>
    """
    raw_text = "1\n05-04-2026\nOfferings for Daily Expense(C)\n516\n💵\nDonation for Holy Mass\n₹8,000.00\n0.00"

    t_text = extract_table_text(raw_text, html)
    assert "05-04-2026" in t_text and "516" in t_text

    entry = {
        "date": "2026-04-05",
        "receipt_number": "516",
        "amount": 8000,
        "head_name": "Offerings for Daily Expense(C)",
        "narration": "Donation for Holy Mass",
    }
    is_valid, reason, details = verify_table_row_evidence(t_text, entry)
    assert is_valid
    assert reason is None
    assert details["voucher"] == "516"
    assert details["amount_match"]
    assert details["date_match"]
    assert details["head_match"]
    assert details["narration_match"]





def test_allowances_five_subheads_exact_and_alias_matching():
    from automation.edayan_dry_run import resolve_subhead_option

    options = {
        "338": "Rev. Fr. Vicar",
        "339": "Rev. Fr. Asst",
        "340": "Accountant",
        "341": "Computer Assistant",
        "342": "Sacristian",
        "343": "Sweeper",
    }

    # Test all 5 Allowances subheads from April Payments sheet
    assert resolve_subhead_option("Rev. Fr. Vicar", options) == ("338", "Rev. Fr. Vicar")
    assert resolve_subhead_option("Accountant", options) == ("340", "Accountant")
    assert resolve_subhead_option("Computer Assistant", options) == ("341", "Computer Assistant")
    assert resolve_subhead_option("Sacristan", options) == ("342", "Sacristian")  # Alias matching
    assert resolve_subhead_option("Sweeper", options) == ("343", "Sweeper")

    # Case-insensitive
    assert resolve_subhead_option("accountant", options) == ("340", "Accountant")
    assert resolve_subhead_option("sacristan", options) == ("342", "Sacristian")


def test_subhead_missing_or_ambiguous_remains_blocked():
    from automation.edayan_dry_run import resolve_subhead_option

    options = {
        "338": "Rev. Fr. Vicar",
        "340": "Accountant",
        "341": "Computer Assistant",
        "342": "Sacristian",
        "343": "Sweeper",
    }

    # Genuinely missing subheads return None (remains blocked)
    assert resolve_subhead_option("Driver", options) is None
    assert resolve_subhead_option("Bishop", options) is None
    assert resolve_subhead_option("", options) is None

    # Ambiguous options return None rather than guessing arbitrarily
    ambiguous_options = {
        "101": "Assistant",
        "102": "Assistant",
    }
    assert resolve_subhead_option("Assistant", ambiguous_options) is None


def test_subhead_delayed_dynamic_loading_and_settling():
    from automation.edayan_entry import wait_for_subhead

    calls = 0

    def mock_read():
        nonlocal calls
        calls += 1
        if calls < 4:
            # First 3 calls: loading/empty options
            return {"state": "loading", "html": "<select></select>", "text": ""}
        # 4th call onwards: settled options
        return {
            "state": "options",
            "html": '<select><option value="0">Select Sub Head</option><option value="338">Rev. Fr. Vicar</option></select>',
            "text": "Rev. Fr. Vicar",
        }

    settled = wait_for_subhead(mock_read, timeout_seconds=5.0, interval_seconds=0.01, stability_checks=2)
    assert settled["state"] == "options"
    assert "Rev. Fr. Vicar" in settled["html"]
    assert calls >= 5  # settled after 2 stability checks


def test_rent_security_direction_aware_mapping_and_multiple_master_matches():
    from automation.edayan_dry_run import _resolve_head

    master = [
        {
            "account_set": "GENERAL",
            "accounting_head_number": "35",
            "extracted_head_name": "Rent Security",
            "human_final_name": "Rent Security Received & Refund",
        },
        {
            "account_set": "GENERAL",
            "accounting_head_number": "88",
            "extracted_head_name": "Rent Security Returned",
            "human_final_name": "Rent Security Received & Refund",
        },
    ]

    crosswalk = [
        {
            "accounting_head_number": "35",
            "transaction_type": "RECEIPT",
            "website_head_value": "157",
            "website_head_label_exact": "Rent Security Received & Refund ( Credit )",
            "mapping_confidence": "0.86",
        },
        {
            "accounting_head_number": "88",
            "transaction_type": "PAYMENT",
            "website_head_value": "157",
            "website_head_label_exact": "Rent Security Received & Refund ( Credit )",
            "mapping_confidence": "1.00",
        },
    ]

    website_options = [
        {"value": "157", "label": "Rent Security Received & Refund ( Credit )", "direction": "RECEIPT"},
    ]

    # PAYMENT direction resolves to Head 88 without ambiguity
    pay_res = _resolve_head("Rent Security Received & Refund", "PAYMENT", master, crosswalk, website_options)
    assert pay_res["status"] == "resolved"
    assert pay_res["website_value"] == "157"
    assert pay_res["website_label"] == "Rent Security Received & Refund ( Credit )"

    # RECEIPT direction resolves to Head 35 without ambiguity
    rec_res = _resolve_head("Rent Security Received & Refund", "RECEIPT", master, crosswalk, website_options)
    assert rec_res["status"] == "resolved"
    assert rec_res["website_value"] == "157"


def test_ambiguous_head_mapping_remains_needs_review():
    from automation.edayan_dry_run import _resolve_head

    master = [
        {
            "account_set": "GENERAL",
            "accounting_head_number": "10",
            "extracted_head_name": "Conflict Head 1",
            "human_final_name": "Conflict Head",
        },
        {
            "account_set": "GENERAL",
            "accounting_head_number": "11",
            "extracted_head_name": "Conflict Head 2",
            "human_final_name": "Conflict Head",
        },
    ]

    # Both mapped to different website values in the same direction -> truly ambiguous
    crosswalk = [
        {
            "accounting_head_number": "10",
            "transaction_type": "PAYMENT",
            "website_head_value": "201",
            "website_head_label_exact": "Conflict 1 ( Debit )",
        },
        {
            "accounting_head_number": "11",
            "transaction_type": "PAYMENT",
            "website_head_value": "202",
            "website_head_label_exact": "Conflict 2 ( Debit )",
        },
    ]

    res = _resolve_head("Conflict Head", "PAYMENT", master, crosswalk, [])
    assert res["status"] == "needs_review"
    assert "Multiple direction-specific crosswalk options" in res["reason"]


def test_voucher_column_match_does_not_match_serial_number_token():
    from automation.edayan_entry import check_live_duplicate, verify_table_row_evidence

    # Table where row serial number is 1, but voucher is 526
    table_text = "1 | 30-04-2026 | | Offerings | 526 | 💵 | Narration | ₹2,266.00 | 0.00 |"

    # Checking for voucher "1" should NOT match row 1's serial number
    entry_v1 = {
        "date": "2026-04-01",
        "receipt_number": "1",
        "amount": 700,
        "head_name": "Miscellaneous Expenses",
        "narration": "Electricity charge",
    }
    dup_status, dup_detail = check_live_duplicate(table_text, entry_v1)
    assert dup_status == "NOT_FOUND"

    is_valid, reason, _ = verify_table_row_evidence(table_text, entry_v1)
    assert not is_valid
    assert "Voucher 1 not found in visible table" in reason

    # Checking for voucher "526" matches voucher column (tokens[4])
    entry_v526 = {
        "date": "2026-04-30",
        "receipt_number": "526",
        "amount": 2266,
        "head_name": "Offerings",
        "narration": "Narration",
    }
    dup_status526, _ = check_live_duplicate(table_text, entry_v526)
    assert dup_status526 == "ALREADY_CONFIRMED_MATCH"


def test_table_row_evidence_matches_website_head_label_and_split_slashes():
    from automation.edayan_entry import verify_table_row_evidence

    table_row = '1 | 16-04-2026 |  | Sunday School Expense | 9 | 💵 | School Maintenance | 0.00 | ₹165.00 | &ensp;'
    entry = {
        'date': '2026-04-16',
        'receipt_number': '9',
        'amount': 165,
        'head_name': 'School/Other Institutions Expenses',
        'narration': 'School Maintenance',
        'head_resolution': {
            'website_label': 'Sunday School Expense ( Debit )',
            'website_value': '137',
        },
    }

    # Matches when website label differs from canonical Excel head name
    is_valid, reason, details = verify_table_row_evidence(table_row, entry)
    assert is_valid is True
    assert reason is None
    assert details['voucher'] == '9'
    assert details['head_match'] is True

    # Rejects when table has an entirely contradictory head
    contradictory_row = '1 | 16-04-2026 |  | Electricity Charges | 9 | 💵 | School Maintenance | 0.00 | ₹165.00 | &ensp;'
    bad_valid, bad_reason, bad_details = verify_table_row_evidence(contradictory_row, entry)
    assert bad_valid is False
    assert bad_details['head_match'] is False
    assert 'Table row mismatch for voucher 9' in bad_reason


def test_non_contiguous_vouchers_remain_eligible_and_no_longest_run_filtering(tmp_path):
    """Rule 2: A voucher gap is not an exclusion reason.
    Non-contiguous voucher numbers (e.g. 1, 2, 514, 515) must all remain eligible without longest-run pruning.
    """
    from automation.batch_manifest import generate_batch_manifest

    wb_path = tmp_path / "test_wb.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Receipts"
    ws.append(["Date", "Receipt Number", "Amount", "Head Name", "Subhead", "Narration"])
    ws.append(["2026-04-03", "1", 725, "Bank Interest", "", "Interest"])
    ws.append(["2026-04-04", "2", 15000, "Bank Withdrawal", "", "Withdrawal"])
    ws.append(["2026-04-05", "514", 1000, "Offerings", "", "Donation"])
    ws.append(["2026-04-06", "515", 2000, "Offerings", "", "Donation"])
    wb.save(wb_path)

    report = {
        "entries": [
            {
                "submission_key": "k1",
                "sheet": "Receipts",
                "row": 2,
                "date": "2026-04-03",
                "receipt_number": "1",
                "amount": 725,
                "head_name": "Bank Interest",
                "subhead": "",
                "narration": "Interest",
                "direction": "RECEIPT",
                "status": "READY",
                "head_resolution": {"website_value": "107"},
            },
            {
                "submission_key": "k2",
                "sheet": "Receipts",
                "row": 3,
                "date": "2026-04-04",
                "receipt_number": "2",
                "amount": 15000,
                "head_name": "Bank Withdrawal",
                "subhead": "",
                "narration": "Withdrawal",
                "direction": "RECEIPT",
                "status": "READY",
                "head_resolution": {"website_value": "4949"},
            },
            {
                "submission_key": "k514",
                "sheet": "Receipts",
                "row": 4,
                "date": "2026-04-05",
                "receipt_number": "514",
                "amount": 1000,
                "head_name": "Offerings",
                "subhead": "",
                "narration": "Donation",
                "direction": "RECEIPT",
                "status": "READY",
                "head_resolution": {"website_value": "128"},
            },
            {
                "submission_key": "k515",
                "sheet": "Receipts",
                "row": 5,
                "date": "2026-04-06",
                "receipt_number": "515",
                "amount": 2000,
                "head_name": "Offerings",
                "subhead": "",
                "narration": "Donation",
                "direction": "RECEIPT",
                "status": "READY",
                "head_resolution": {"website_value": "128"},
            },
        ]
    }
    journal = {"entries": {}}

    manifest = generate_batch_manifest(report, journal, wb_path, direction="RECEIPT")
    auth_vouchers = [e["receipt_number"] for e in manifest.authorized_entries]
    assert len(manifest.authorized_entries) == 4
    assert auth_vouchers == ["1", "2", "514", "515"]
    assert len(manifest.excluded_entries) == 0


def test_full_source_coverage_accounts_for_every_source_row(tmp_path):
    """Rule 1 & Rule 6: Source transaction set = Confirmed + Authorized + Explicit exclusions.
    Every source row must be accounted for without silent loss.
    """
    from automation.batch_manifest import generate_batch_manifest, check_source_completeness

    wb_path = tmp_path / "test_wb.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Receipts"
    ws.append(["Date", "Receipt Number", "Amount", "Head Name", "Subhead", "Narration"])
    ws.append(["2026-04-03", "1", 725, "Bank Interest", "", "Interest"])
    ws.append(["2026-04-04", "2", 15000, "Bank Withdrawal", "", "Withdrawal"])
    wb.save(wb_path)

    report = {
        "entries": [
            {
                "submission_key": "k1",
                "sheet": "Receipts",
                "row": 2,
                "date": "2026-04-03",
                "receipt_number": "1",
                "amount": 725,
                "head_name": "Bank Interest",
                "subhead": "",
                "narration": "Interest",
                "direction": "RECEIPT",
                "status": "READY",
                "head_resolution": {"website_value": "107"},
            },
            {
                "submission_key": "k2",
                "sheet": "Receipts",
                "row": 3,
                "date": "2026-04-04",
                "receipt_number": "2",
                "amount": 15000,
                "head_name": "Bank Withdrawal",
                "subhead": "",
                "narration": "Withdrawal",
                "direction": "RECEIPT",
                "status": "READY",
                "head_resolution": {"website_value": "4949"},
            },
        ]
    }
    journal = {
        "entries": {
            "k1": {"status": "confirmed", "evidence": {"record_id": "1"}}
        }
    }

    manifest = generate_batch_manifest(report, journal, wb_path, direction="RECEIPT")
    completeness = check_source_completeness(report, manifest, journal)
    assert completeness["complete"] is True
    assert completeness["source_count"] == 2
    assert completeness["confirmed_count"] == 1
    assert completeness["authorized_count"] == 1
    assert len(completeness["unaccounted_entries"]) == 0
    assert len(manifest.authorized_entries) == 1
    assert manifest.authorized_entries[0]["receipt_number"] == "2"


def test_approved_workbook_authority_and_kappad_option_4949():
    """Rule 3, 4, 6 & 7: The approved Excel workbook is the single source of truth.
    Kappad matches option 4949 with '(Debit)' suffix without altering approved head or mapping to Pala.
    """
    from automation.edayan_dry_run import run_dry_run

    journal_path = Path("data/archive/2026-04/execution_logs/eday_entry_progress.json")
    if not journal_path.exists():
        journal_path = Path("data/output/eday_entry_progress.json")

    report = run_dry_run(
        Path("data/monthly/2026-04/2026-04_ready_for_entry.xlsx"),
        Path("data/reference/general_head_crosswalk.csv"),
        journal_path,
        master_path=Path("data/reference/active/accounting_heads_general_verified.csv"),
        website_html_path=Path("data/reference/website_heads.csv"),
        subheads_path=Path("data/reference/website_subheads_observed.json"),
    )

    # Find row 21 (Voucher 2, Kappad)
    kappad_entries = [
        e for e in report["entries"]
        if e.get("sheet") == "Receipts" and str(e.get("receipt_number")) == "2"
    ]
    assert len(kappad_entries) == 1
    k_entry = kappad_entries[0]
    assert k_entry["head_name"] == "The South Indian Bank Kappad"
    assert k_entry["status"] == "READY"
    assert k_entry["head_resolution"]["website_value"] == "4949"
    assert k_entry["head_resolution"]["website_label"] == "The South Indian Bank Kappad ( Debit )"
    assert "Pala" not in k_entry["head_resolution"]["website_label"]
    assert k_entry["narration"] == "South Indian Bank S.B. withdraw"

    # Find row 3 (Voucher 1, Bank Interest)
    interest_entries = [
        e for e in report["entries"]
        if e.get("sheet") == "Receipts" and str(e.get("receipt_number")) == "1"
    ]
    assert len(interest_entries) == 1
    i_entry = interest_entries[0]
    assert i_entry["head_name"] == "Bank Interest/ Other Interest Received (C)"
    assert i_entry["status"] == "READY"
    assert i_entry["head_resolution"]["website_value"] == "107"
    assert i_entry["narration"] == "In South Indian Bank (1-1-26/31-3-26)"
    assert i_entry["subhead"] == ""
