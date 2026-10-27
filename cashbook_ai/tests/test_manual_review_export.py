import json
from pathlib import Path

import openpyxl

from src.export.manual_review_exporter import MANUAL_REVIEW_COLUMNS, ManualReviewExporter
from generate_april_manual_review import reconcile_rows
from src.models import (
    AccountSet,
    AccountingHead,
    ProcessedRow,
    RecordType,
    TransactionType,
)
from src.validation.head_master import HeadMasterRepository


def make_row(**values):
    defaults = {
        "source_image": "page_001.jpg",
        "source_page": "1",
        "source_row": 1,
        "resolved_date": "2026-04-01",
        "raw_voucher_number": "531-535",
        "accounting_head_name": "Canonical Head",
        "subhead": None,
        "narration_english": "Narration",
        "record_type": RecordType.TRANSACTION,
        "transaction_type": TransactionType.PAYMENT,
        "payment_amount": 25.0,
    }
    defaults.update(values)
    return ProcessedRow(**defaults)


def test_manual_review_workbook_has_exact_sheets_and_columns(tmp_path: Path):
    output = tmp_path / "manual_review.xlsx"
    ManualReviewExporter().export([make_row()], output)

    workbook = openpyxl.load_workbook(output, read_only=True, data_only=True)
    assert workbook.sheetnames == ["Receipts", "Payments"]
    for worksheet in workbook.worksheets:
        assert [cell.value for cell in worksheet[1]] == MANUAL_REVIEW_COLUMNS
        assert worksheet.max_column == 6


def test_manual_review_separates_directions_and_excludes_balances(tmp_path: Path):
    output = tmp_path / "manual_review.xlsx"
    rows = [
        make_row(transaction_type=TransactionType.RECEIPT, receipt_amount=10.0, payment_amount=None),
        make_row(transaction_type=TransactionType.PAYMENT, receipt_amount=None, payment_amount=20.0),
        make_row(record_type=RecordType.OPENING_BALANCE, transaction_type=TransactionType.INFORMATIONAL),
        make_row(record_type=RecordType.CLOSING_BALANCE, transaction_type=TransactionType.INFORMATIONAL),
    ]
    ManualReviewExporter().export(rows, output)

    workbook = openpyxl.load_workbook(output, read_only=True, data_only=True)
    assert workbook["Receipts"].max_row == 2
    assert workbook["Payments"].max_row == 2
    assert workbook["Receipts"].cell(2, 3).value == 10
    assert workbook["Payments"].cell(2, 3).value == 20


def test_reconciled_master_uses_human_final_name_and_preserves_subhead_source(tmp_path: Path):
    master = tmp_path / "master.csv"
    master.write_text(
        "account_set,accounting_head_number,extracted_head_name,human_final_name,review_status,human_comment,accounting_head_name,source_type\n"
        "GENERAL,52,Allowances: Original Subhead,Allowances,CORRECTED,,Allowances,MANUAL_VERIFIED\n",
        encoding="utf-8",
    )
    repository = HeadMasterRepository(master)
    assert repository.get_head_name(AccountSet.GENERAL, "52") == "Allowances"
    assert repository.get_extracted_head_name(AccountSet.GENERAL, "52") == "Allowances: Original Subhead"
    assert repository.get_head_name(AccountSet.GENERAL, "999") is None


def test_manual_review_preserves_unusual_receipt_number_and_blank_subhead(tmp_path: Path):
    output = tmp_path / "manual_review.xlsx"
    ManualReviewExporter().export([
        make_row(raw_voucher_number="531-535", subhead=None)
    ], output)
    workbook = openpyxl.load_workbook(output, read_only=True, data_only=True)
    row = next(workbook["Payments"].iter_rows(min_row=2, max_row=2, values_only=True))
    assert row[1] == "531-535"
    assert row[4] in ("", None)


def test_source_identifiers_reconcile_without_dropped_rows():
    raw = [
        make_row(source_row=1),
        make_row(source_row=2),
    ]
    processed = [
        make_row(source_row=1),
        make_row(source_row=2),
    ]
    assert reconcile_rows(raw, processed) == ([], [], [], [])
