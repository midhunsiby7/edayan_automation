"""Recover the saved April review decisions and produce a corrected workbook."""

from __future__ import annotations

import argparse
import json
import shutil
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl

from generate_physical_column_review import load_processed_rows
from src.export.manual_review_exporter import ManualReviewExporter
from src.models import RecordType, TransactionType
from src.review.physical_column_review import extraction_sha256, source_key


ROOT = Path(__file__).parent
RAW_PATH = ROOT / "data/monthly/raw_extractions.json"
REVIEW_PATH = ROOT / "data/monthly/2026-04/2026-04_physical_column_review.xlsx"
OUTPUT_PATH = ROOT / "data/monthly/2026-04/2026-04_corrected.xlsx"
CORRECTIONS_PATH = ROOT / "data/monthly/2026-04/physical_column_corrections.json"
RECONCILIATION_PATH = ROOT / "data/monthly/2026-04/april_corrected_reconciliation.json"
EXCEPTIONS_PATH = ROOT / "data/monthly/2026-04/april_corrected_exceptions.json"


def normalize_direction(value: Any) -> str | None:
    if value is None:
        return None
    normalized = str(value).strip().casefold()
    if normalized in {"receipt", "receipts"}:
        return "Receipt"
    if normalized in {"payment", "payments"}:
        return "Payment"
    if normalized in {"ambiguous", "unclear"}:
        return "Ambiguous"
    if not normalized:
        return None
    return None


def backup_review_workbook(path: Path) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = path.with_name(f"{path.stem}_user_saved_backup_{timestamp}{path.suffix}")
    shutil.copy2(path, backup)
    return backup


def read_review_decisions(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    workbook = openpyxl.load_workbook(path, data_only=True)
    if "Physical_Column_Review" not in workbook.sheetnames:
        raise ValueError("Saved review workbook is missing Physical_Column_Review")
    sheet = workbook["Physical_Column_Review"]
    headers = [cell.value for cell in sheet[1]]
    required = {
        "Source Image",
        "Source Page",
        "Source Row",
        "Record Type",
        "Original Model Classification",
        "Human-Verified Direction",
        "Review Status",
        "Correction Reason",
    }
    if not required.issubset(headers):
        raise ValueError("Saved review workbook does not contain the required decision columns")
    positions = {header: index for index, header in enumerate(headers)}
    records: list[dict[str, Any]] = []
    seen: set[tuple[str, str, int]] = set()
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if all(value is None for value in values):
            continue
        key = (
            str(values[positions["Source Image"]]),
            str(values[positions["Source Page"]]),
            int(values[positions["Source Row"]]),
        )
        if key in seen:
            raise ValueError(f"Duplicate source identifier in saved workbook: {key}")
        seen.add(key)
        records.append(
            {
                "source_image": key[0],
                "source_page": key[1],
                "source_row": key[2],
                "record_type": str(values[positions["Record Type"]]),
                "model_classification": str(values[positions["Original Model Classification"]]),
                "human_direction_raw": values[positions["Human-Verified Direction"]],
                "human_direction": normalize_direction(values[positions["Human-Verified Direction"]]),
                "review_status": str(values[positions["Review Status"]] or ""),
                "correction_reason": str(values[positions["Correction Reason"]] or ""),
            }
        )
    metadata = {}
    if "Review_Metadata" in workbook.sheetnames:
        metadata = {
            str(row[0].value): row[1].value
            for row in workbook["Review_Metadata"].iter_rows(min_row=1, max_col=2)
            if row[0].value is not None
        }
    return records, metadata


def _direction_for_record(record: dict[str, Any], processed: Any) -> tuple[str | None, str, str]:
    explicit = record["human_direction"]
    if explicit in {"Receipt", "Payment"}:
        return explicit, "USER_SAVED_REVIEW", "Explicit direction recovered from the saved review workbook"
    if processed.record_type != RecordType.TRANSACTION:
        return None, "INFORMATIONAL", "Informational balance excluded from transaction export"
    if explicit == "Ambiguous":
        return None, "AMBIGUOUS", "User marked the direction ambiguous"
    return None, "UNRESOLVED", "No valid explicit direction was saved"


def process(
    review_path: Path = REVIEW_PATH,
    output_path: Path = OUTPUT_PATH,
    corrections_path: Path = CORRECTIONS_PATH,
    reconciliation_path: Path = RECONCILIATION_PATH,
    exceptions_path: Path = EXCEPTIONS_PATH,
) -> dict[str, Any]:
    backup = backup_review_workbook(review_path)
    raw_rows, processed_rows = load_processed_rows()
    review_records, review_metadata = read_review_decisions(review_path)
    processed_by_key = {source_key(row): row for row in processed_rows}
    raw_by_key = {source_key(row): row for row in raw_rows}
    if len(processed_by_key) != len(processed_rows) or len(raw_by_key) != len(raw_rows):
        raise ValueError("Duplicate source identifiers detected in raw or processed records")
    review_by_key = {
        (record["source_image"], record["source_page"], record["source_row"]): record
        for record in review_records
    }
    missing_review = sorted(set(processed_by_key) - set(review_by_key))
    extra_review = sorted(set(review_by_key) - set(processed_by_key))
    if missing_review or extra_review:
        raise ValueError(f"Review/source mismatch: missing={missing_review}, extra={extra_review}")

    decisions: list[dict[str, Any]] = []
    directions: dict[tuple[str, str, int], TransactionType] = {}
    exceptions: list[dict[str, Any]] = []
    user_edit_count = 0
    for key, processed in sorted(processed_by_key.items()):
        review = review_by_key[key]
        direction, source, reason = _direction_for_record(review, processed)
        model = processed.transaction_type.value
        if direction and direction != ("Receipt" if model == "RECEIPT" else "Payment" if model == "PAYMENT" else None):
            user_edit_count += 1
            reason += "; differs from original model classification"
        if direction:
            directions[key] = TransactionType.RECEIPT if direction == "Receipt" else TransactionType.PAYMENT
        elif processed.record_type == RecordType.TRANSACTION:
            exceptions.append(
                {
                    "source_image": key[0],
                    "source_page": key[1],
                    "source_row": key[2],
                    "reason": reason,
                    "model_classification": model,
                    "saved_direction": review["human_direction_raw"],
                }
            )
        decisions.append(
            {
                "source_image": key[0],
                "source_page": key[1],
                "source_row": key[2],
                "original_model_classification": model,
                "original_model_detected_amount_column": processed.detected_amount_column,
                "selected_direction": direction,
                "correction_source": source,
                "verification_method": "USER_SAVED_REVIEW_WORKBOOK" if source == "USER_SAVED_REVIEW" else source,
                "correction_reason": review["correction_reason"] or reason,
                "review_status": review["review_status"],
                "head_number": processed.accounting_head_number,
                "head_name": processed.accounting_head_name,
                "source_extraction_sha256": extraction_sha256(RAW_PATH),
            }
        )

    head_exceptions = [
        {
            "source_image": key[0],
            "source_page": key[1],
            "source_row": key[2],
            "raw_head_number": raw_by_key[key].raw_head_number,
            "normalized_head_number": processed_by_key[key].accounting_head_number,
            "head_name": processed_by_key[key].accounting_head_name,
            "reason": "Head number is absent from the active verified General master",
        }
        for key in processed_by_key
        if processed_by_key[key].record_type == RecordType.TRANSACTION
        and processed_by_key[key].accounting_head_name is None
    ]
    if exceptions:
        output_path = output_path.with_name(output_path.stem + "_draft" + output_path.suffix)
    ManualReviewExporter().export_with_directions(processed_rows, directions, output_path)

    payload = {
        "schema_version": "april-physical-correction-1.0",
        "source_path": str(RAW_PATH),
        "source_sha256": extraction_sha256(RAW_PATH),
        "review_workbook_path": str(review_path),
        "review_workbook_backup_path": str(backup),
        "review_workbook_metadata": review_metadata,
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "raw_record_count": len(raw_rows),
        "processed_record_count": len(processed_rows),
        "eligible_transaction_count": sum(row.record_type == RecordType.TRANSACTION for row in processed_rows),
        "informational_count": sum(row.record_type != RecordType.TRANSACTION for row in processed_rows),
        "user_direction_corrections": user_edit_count,
        "records": decisions,
    }
    corrections_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    exceptions_path.write_text(
        json.dumps(
            {
                "source_sha256": payload["source_sha256"],
                "direction_exceptions": exceptions,
                "head_exceptions": head_exceptions,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    receipts = [key for key, direction in directions.items() if direction == TransactionType.RECEIPT]
    payments = [key for key, direction in directions.items() if direction == TransactionType.PAYMENT]
    reconciliation = {
        "source_record_count": len(raw_rows),
        "processed_record_count": len(processed_rows),
        "eligible_transaction_count": payload["eligible_transaction_count"],
        "informational_excluded_count": payload["informational_count"],
        "user_direction_corrections": user_edit_count,
        "receipt_count": len(receipts),
        "payment_count": len(payments),
        "exported_count": len(directions),
        "unresolved_direction_count": len(exceptions),
        "head_exception_count": len(head_exceptions),
        "exported_source_keys": [f"{k[0]}|{k[1]}|{k[2]}" for k in receipts + payments],
        "direction_exception_source_keys": [f"{e['source_image']}|{e['source_page']}|{e['source_row']}" for e in exceptions],
        "head_exception_source_keys": [f"{e['source_image']}|{e['source_page']}|{e['source_row']}" for e in head_exceptions],
        "final_workbook": str(output_path) if not exceptions else None,
        "draft_workbook": str(output_path) if exceptions else None,
        "fully_direction_reconciled": not exceptions,
        "head_names_all_resolved": not head_exceptions,
    }
    reconciliation_path.write_text(json.dumps(reconciliation, ensure_ascii=False, indent=2), encoding="utf-8")
    return {
        "backup": str(backup),
        "output": str(output_path) if output_path.exists() else None,
        "corrections": str(corrections_path),
        "reconciliation": str(reconciliation_path),
        "exceptions": str(exceptions_path),
        **reconciliation,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Recover saved April review decisions and generate corrected workbook.")
    parser.add_argument("--review", type=Path, default=REVIEW_PATH)
    args = parser.parse_args()
    print(json.dumps(process(review_path=args.review), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
